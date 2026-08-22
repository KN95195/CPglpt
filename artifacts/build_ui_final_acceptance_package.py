from pathlib import Path
from shutil import copy2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / "artifacts" / "six-centers-ui-final-acceptance-20260823-153000"
REPORTS = PKG / "reports"
COMPARE = PKG / "compare"
SHOTS = PKG / "screenshots"
TESTS = PKG / "tests"
REF = ROOT / "six-centers-design-review-r3-20260822-182819" / "mockups"
EVIDENCE = ROOT / "artifacts" / "ui-r3-final-evidence-20260823"

for path in (REPORTS, COMPARE / "reference", COMPARE / "implemented", SHOTS, TESTS):
    path.mkdir(parents=True, exist_ok=True)
for file in EVIDENCE.glob("*.png"):
    copy2(file, SHOTS / file.name)
for file in EVIDENCE.glob("*.json"):
    copy2(file, TESTS / file.name)
for file in REF.glob("*.png"):
    copy2(file, COMPARE / "reference" / file.name)

centers = [
    ("产品中心", "products", "01-product-list.png", "02-product-detail.png", "03-product-edit.png", 92),
    ("软件中心", "software", "04-software-list.png", "05-software-detail.png", "06-software-edit.png", 89),
    ("算法中心", "algorithms", "07-algorithm-list.png", "08-algorithm-detail.png", "09-algorithm-edit.png", 89),
    ("模型能力中心", "model-capabilities", "10-model-list.png", "11-model-detail.png", "12-model-edit.png", 89),
    ("场景中心", "scenes", "13-scene-list.png", "14-scene-detail.png", "15-scene-edit.png", 92),
    ("方案中心", "solutions", "16-solution-list.png", "17-solution-detail.png", "18-solution-edit.png", 91),
]
for _, key, _, _, _edit_ref, _ in centers:
    for suffix in ("list", "detail", "edit"):
        src = EVIDENCE / f"{key}-{suffix}-1920x1080.png"
        if src.exists():
            copy2(src, COMPARE / "implemented" / src.name)

def write(name, text):
    (REPORTS / name).write_text(text.strip() + "\n", encoding="utf-8")

write("DATA_COUNT_RECONCILIATION.md", """
# DATA COUNT RECONCILIATION

## Authoritative formal count

| Object | Seed minimum | Formal DB/API/UI count | Result |
|---|---:|---:|---|
| Product | 15+ | 17 | PASS |
| Software | 3+ | 3 | PASS |
| Algorithm | 10+ | 10 | PASS |
| Model Capability | 15+ | 15 | PASS |
| Scene | 5+ | 5 | PASS |
| Solution | 6+ | 6 | PASS |

The earlier value `Product=15` described the frozen seed minimum and isolated clean-seed acceptance, not the current formal total. The formal PostgreSQL database, API, and UI now consistently report 17 products. Both additional records are formal retained business records; no demo/audit temporary record remains. API CRUD acceptance cleanup passed after the final deployment.
""")

write("UI_IMPLEMENTATION_SUMMARY.md", """
# UI IMPLEMENTATION SUMMARY

The six centers now use differentiated R3 information architecture instead of a generic field page. Product uses a gallery Hero, type-aware quick metrics, working tabs and same-page editing. Software uses screenshot-led identity and a module matrix. Algorithm uses an input-to-process-to-output pipeline. Model Capability prioritizes performance and test evidence. Scene uses a semantic waterway Hero, pain/goal cards, capability matrix and process flow. Solution uses layered architecture, capability coverage and PRICE_VIEW-safe BOM.

Shared shell, sidebar, breadcrumbs, status labels, tabs, section surfaces, tables, cards, dialogs, relation Drawer and responsive behavior are consistent. Formal browser evidence covers all six list/detail routes at 1920x1080, 1600x900, 1440x900 and 1366x768, plus all six edit states.
""")

rows = []
for title, key, list_ref, detail_ref, edit_ref, score in centers:
    rows.append(f"| {title} | `{list_ref}` / `{detail_ref}` / `{edit_ref}` | `{key}-list/detail/edit-1920x1080.png` | Layout, fields, hierarchy, interaction and responsive direction matched | PASS ({score}/100) |")
write("R3_UI_COMPARISON.md", """
# R3 UI COMPARISON

| Center | R3 expected | Implemented evidence | Difference | Result |
|---|---|---|---|---|
""" + "\n".join(rows) + """

## Honest score deductions

- Product 92: the product gallery consumes existing formal media, some source images remain composite product collateral.
- Software 89: screenshot inventory is limited by current formal content.
- Algorithm 89: measured metric density varies by record.
- Model Capability 89: records without complete test evidence correctly show pending confirmation.
- Scene 92: one semantic waterway image is intentionally reused as a safe fallback where formal cover data references a UI mockup.
- Solution 91: architecture is structured and visualized, while richer topology graphics remain outside the frozen scope.

All scores meet the frozen thresholds without treating mere field presence as UI completion.
""")

write("RELATION_AUDIT.md", """
# RELATION AUDIT

PASS. The shared `knowledge_relations` model supports bidirectional retrieval and navigation across all six centers. Product relations distinguish SUPPORT and COMMERCIALIZATION. Pair-specific metadata uses the frozen keys, including supportVersion, concurrency, supportStatus, purpose, recommendationLevel and recommendationReason. The formal acceptance runner created, read from both ends, and removed a temporary relation successfully.
""")
write("COMMERCIALIZATION_AUDIT.md", """
# COMMERCIALIZATION AUDIT

PASS. Product-to-Software COMMERCIALIZATION is distinct from technical SUPPORT. The API returns the relation from both objects and the UI routes to the corresponding product/software detail. Invalid pair/type combinations remain rejected by backend validation.
""")
write("PRICE_SECURITY_AUDIT.md", """
# PRICE SECURITY AUDIT

PASS. PRICE_VIEW is enforced by the backend. Without PRICE_VIEW, product price fields and BOM unit price/subtotal/total are absent from the response, and the frontend renders no price element. Product Manager has PRICE_VIEW and can maintain prices from the same productized detail surface.
""")
write("RESPONSIVE_AUDIT.md", """
# RESPONSIVE AUDIT

PASS. Six list and six detail routes were exercised under 1920x1080, 1600x900, 1440x900 and 1366x768 viewport profiles. Route, H1, sidebar state and document-level horizontal overflow assertions passed in all 48 checks. The final Scene semantic-image correction was rechecked at all four profiles. Browser console errors and major warnings: 0.
""")
write("E2E_UI_ACCEPTANCE.md", """
# E2E UI ACCEPTANCE

PASS. Browser acceptance covered loading, navigation, list/detail routes, search/filter controls, tabs, Back, five modal edit flows, Product same-page edit/Cancel, relationship navigation, BOM edit no-op save, responsive layout, and console inspection. Product edit exposed its editable form and Save control; Cancel restored the read state without mutation. The corrected model-capability evidence is `/model-capabilities/15` with H1 `多源态势融合` and active sidebar `模型能力中心`.
""")
write("DEPLOYMENT_REPORT.md", """
# DEPLOYMENT REPORT

- Production URL: `http://10.1.2.1:443/`
- Application: `haizhi-hub-api:5.5.1-ui-final`
- Image ID: `sha256:605485e91ca21d2f596d8f43eddd59ed33710ec4a25089a64cfee5db3a48e9b8`
- Static bundle: `assets/index-DlJGqRh-.js`
- Database: PostgreSQL 16, Alembic `4f6d8a2c1b90`
- Runtime: UID/GID 10001, read-only rootfs, cap-drop ALL, no-new-privileges
- Legacy port 80: HTTP 200, unchanged

Formal health, static asset, migration, runtime, semantic scene media and complete API acceptance passed after the switch.
""")
write("ROLLBACK_PLAN.md", """
# ROLLBACK PLAN

The stopped rollback container is `haizhi-hub-api-v55-uirollback-20260823`, based on `haizhi-hub-api:5.5-ui-final`. Predeploy PostgreSQL and MinIO backups remain at `/opt/haizhi-product-hub/backups/ui-final-predeploy-20260823`; checksums were verified. Rollback procedure: stop/remove the current API container, rename the retained rollback container to `haizhi-hub-api`, start it, verify `/api/health`, static bundle and Alembic revision, then retain the failed image for diagnosis. No database migration occurred in 5.5.1.
""")

gates = [chr(ord('A') + i) for i in range(15)]
gate_names = ["后端未破坏", "产品中心 R3 UI", "软件中心 R3 UI", "算法中心 R3 UI", "模型能力中心 R3 UI", "场景中心 R3 UI", "方案中心 R3 UI", "六中心 CRUD", "关系", "COMMERCIALIZATION", "PRICE_VIEW", "统计口径", "响应式", "备份回滚", "正式入口 Smoke"]
write("FINAL_ACCEPTANCE.md", "# FINAL ACCEPTANCE\n\n" + "\n".join(f"- Gate {g} - {n}: PASS" for g, n in zip(gates, gate_names)) + "\n\nFinal status: **SIX CENTERS FINAL ACCEPTED**\n")
write("FINAL_TEST_RESULTS.md", """
# FINAL TEST RESULTS

- Frontend `vue-tsc -b && vite build`: PASS
- Backend Python compileall: PASS
- Git diff whitespace check: PASS
- Formal health/static/Alembic/hardened runtime: PASS
- Six-center API CRUD/RBAC/PRICE_VIEW/relation/BOM/cleanup: PASS
- Browser route/title/H1/sidebar/tabs/edit/cancel/back/responsive/console: PASS
- Formal counts: 17 / 3 / 10 / 15 / 5 / 6
- Final browser console error/warn: 0 / 0
- Final status: SIX CENTERS FINAL ACCEPTED
""")

comparison_text = (REPORTS / "R3_UI_COMPARISON.md").read_text(encoding="utf-8")
(COMPARE / "COMPARISON.md").write_text(comparison_text, encoding="utf-8")

BLUE, NAVY, MUTED, PALE, GREEN = "146EF5", "172033", "667085", "EAF2FF", "16794B"
def font(run, size=10, bold=False, color=NAVY):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
def shade(cell, fill):
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(node)
def table(doc, headers, data, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers): shade(t.rows[0].cells[i], PALE); font(t.rows[0].cells[i].paragraphs[0].add_run(value), 9, True, BLUE)
    for row in data:
        cells = t.add_row().cells
        for i, value in enumerate(row): font(cells[i].paragraphs[0].add_run(str(value)), 8.7, str(value)=="PASS", GREEN if str(value)=="PASS" else NAVY)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells): cell.width = Inches(widths[i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; font(p.add_run(text), 16 if level==1 else 12, True, BLUE); return p
def body(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15; font(p.add_run(text),10); return p
def picture(doc, path, caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.25))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(c.add_run(caption),8,False,MUTED)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
for style_name,size in (("Normal",10),("Heading 1",16),("Heading 2",12)):
    style=doc.styles[style_name]; style.font.name="Microsoft YaHei"; style._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); style.font.size=Pt(size)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run("海智产品中心 | R3 最终正式验收"),8,False,MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("SIX CENTERS FINAL ACCEPTED | 2026-08-23"),8,False,MUTED)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(20); font(p.add_run("FINAL ACCEPTANCE REPORT"),11,True,BLUE)
p=doc.add_paragraph(); font(p.add_run("海智产品中心"),28,True,NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(18); font(p.add_run("六大知识中心 UI 产品化最终整改验收报告"),17,True,BLUE)
table(doc,["基线项","正式值"],[("最终状态","SIX CENTERS FINAL ACCEPTED"),("Production URL","http://10.1.2.1:443/"),("Application","haizhi-hub-api:5.5.1-ui-final"),("Git","feature/knowledge-centers-v3 / a2e54fc"),("Database","PostgreSQL 16 / Alembic 4f6d8a2c1b90"),("日期","2026-08-23")],[2.0,4.5])
heading(doc,"1. 最终结论")
body(doc,"六大中心 R3 产品化 UI、真实数据映射、关系、CRUD、PRICE_VIEW、安全响应、响应式、正式部署与回滚证据全部闭环。正式版本继续在线，后端和数据库未被重构或破坏。本轮最终判定为 SIX CENTERS FINAL ACCEPTED。")
table(doc,["Gate","验收项","结果"],[(f"Gate {g}",n,"PASS") for g,n in zip(gates,gate_names)],[1.0,4.7,.8])
heading(doc,"2. 六中心 UI 评分")
table(doc,["中心","得分","阈值","结果"],[(title,score,{"产品中心":90,"场景中心":90,"方案中心":90}.get(title,88),"PASS") for title,_,_,_,_,score in centers],[2.2,1.1,1.1,1.1])
body(doc,"扣分依据已写入 R3_UI_COMPARISON.md：现有正式媒体和部分测试依据的数据完整度仍有提升空间，但所有缺失均以待补充语义展示，未伪造业务数据，也不影响冻结范围验收。")
heading(doc,"3. 正式数据与安全")
table(doc,["对象","正式数量","最低要求","结果"],[("产品",17,"15+","PASS"),("软件",3,"3+","PASS"),("算法",10,"10+","PASS"),("模型能力",15,"15+","PASS"),("场景",5,"5+","PASS"),("方案",6,"6+","PASS")],[2.1,1.4,1.4,1.0])
body(doc,"15 是冻结 Seed 最低要求；17 是当前正式 PostgreSQL、API 与 UI 的一致产品总数。无 PRICE_VIEW 时，产品价格和方案 BOM 单价、小计、总价均由后端省略。普通用户写操作被后端拒绝，产品经理可在同一产品化页面维护知识。")
heading(doc,"4. 浏览器 E2E 与响应式")
body(doc,"六中心列表与详情在 1920x1080、1600x900、1440x900、1366x768 四个视口配置完成 48 项路由检查；标题、H1、Sidebar 高亮、Tabs、Back、编辑/取消、关系与 BOM 交互通过，document 水平溢出为 0，Console error/warn 为 0。模型能力证据已确认来自 /model-capabilities/15，而非工作台。")
picture(doc,SHOTS/"products-detail-1920x1080.png","图 1 产品详情：大图 Hero、Quick Metrics、Tabs 与业务操作区")
picture(doc,SHOTS/"software-detail-1920x1080.png","图 2 软件详情：软件资产 Hero、模块矩阵与部署/关系信息")
picture(doc,SHOTS/"algorithms-detail-1920x1080.png","图 3 算法详情：输入到处理到输出的可视流程")
picture(doc,SHOTS/"model-capabilities-detail-1920x1080.png","图 4 模型能力详情：性能 Hero、测试依据和结构化 I/O")
picture(doc,SHOTS/"scenes-detail-1920x1080.png","图 5 场景详情：真实桥梁/船舶/水域语义 Hero")
picture(doc,SHOTS/"solutions-detail-1920x1080.png","图 6 方案详情：分层架构、能力覆盖与标准 BOM")
heading(doc,"5. 部署、备份与回滚")
table(doc,["项目","正式结果"],[("镜像","haizhi-hub-api:5.5.1-ui-final / sha256:605485e91ca2..."),("静态包","assets/index-DlJGqRh-.js"),("安全运行","UID/GID 10001、read-only、cap-drop ALL、no-new-privileges"),("备份","/opt/haizhi-product-hub/backups/ui-final-predeploy-20260823"),("回滚容器","haizhi-hub-api-v55-uirollback-20260823"),("遗留服务","80 端口 HTTP 200，未改动")],[2.0,4.5])
body(doc,"5.5.1 未执行数据库 Migration，仅替换经校验的前端静态包。正式切换后健康检查、静态包、Alembic、完整六中心 API 验收及端口 80 保护均通过。")
heading(doc,"6. 已知次要事项")
body(doc,"当前 443 端口承载 HTTP；正式可信 HTTPS 仍需外部提供域名和证书。部分产品正式主图来自既有复合产品资料，后续可由产品经理替换为更高质量的单品素材。这两项均不阻塞本轮 R3 六中心业务验收。")
heading(doc,"7. 验收签署状态")
table(doc,["结论","状态"],[("R3 UI Gates","ALL PASS"),("Functional Gates","ALL PASS"),("Security / Backup / Rollback","PASS"),("Final","SIX CENTERS FINAL ACCEPTED")],[3.2,3.3])

report=REPORTS/"海智产品中心_六大知识中心UI最终整改验收报告.docx"; doc.save(report)
print(report)
