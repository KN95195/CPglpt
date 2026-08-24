from pathlib import Path
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[3]
ARTIFACT = ROOT / "artifacts" / "phase2-business-implementation-20260824-final"
REPORTS = ARTIFACT / "reports"
SCREENSHOTS = ARTIFACT / "screenshots"
OUTPUT = REPORTS / "海智产品中心_Phase2业务实施验收报告.docx"

BLUE = "1769AA"
BLUE_DARK = "0B3E66"
BLUE_LIGHT = "EAF4FB"
GRAY = "F2F4F7"
INK = "1F2937"
MUTED = "667085"
GREEN = "16794A"
AMBER = "9A6700"
RED = "B42318"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def font_run(run, size=11, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    font_run(run, size=9, color=MUTED)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Title": (24, BLUE_DARK, 0, 6),
        "Subtitle": (11, MUTED, 0, 12),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, BLUE_DARK, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font_run(p.add_run("海智产品中心  |  Phase 2 业务实施验收"), size=9, color=MUTED)
    footer = section.footer
    ftable = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(ftable, [6200, 3160])
    ftable.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    font_run(ftable.cell(0, 0).paragraphs[0].add_run("内部验收材料  |  2026-08-24"), size=9, color=MUTED)
    add_page_number(ftable.cell(0, 1).paragraphs[0])


def add_title_block(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font_run(p.add_run("海智产品中心"), size=24, bold=True, color=BLUE_DARK)
    p2 = doc.add_paragraph(style="Subtitle")
    font_run(p2.add_run("Phase 2 业务实施验收报告"), size=16, bold=True, color=BLUE)
    rule = doc.add_table(rows=1, cols=1)
    set_table_geometry(rule, [9360])
    rule.rows[0].height = Inches(0.06)
    set_cell_shading(rule.cell(0, 0), BLUE)
    rule.cell(0, 0).text = ""

    meta = doc.add_table(rows=5, cols=2)
    set_table_geometry(meta, [2700, 6660])
    fields = [
        ("报告状态", "PHASE 2 BUSINESS IMPLEMENTATION NOT READY"),
        ("正式环境", "http://10.1.2.1:443/"),
        ("应用版本", "haizhi-hub-api:6.0.5-phase2-style-fix"),
        ("代码与数据", "Git 1c91a3a  |  Alembic c3d4e5f60718"),
        ("报告日期", "2026-08-24"),
    ]
    for i, (label, value) in enumerate(fields):
        set_cell_shading(meta.cell(i, 0), BLUE_LIGHT)
        font_run(meta.cell(i, 0).paragraphs[0].add_run(label), size=10, bold=True, color=BLUE_DARK)
        color = RED if i == 0 else INK
        font_run(meta.cell(i, 1).paragraphs[0].add_run(value), size=10, bold=i == 0, color=color)
    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360])
    set_cell_shading(callout.cell(0, 0), "FFF4E5")
    p = callout.cell(0, 0).paragraphs[0]
    font_run(p.add_run("验收结论："), size=11, bold=True, color=AMBER)
    font_run(p.add_run("平台侧可执行开发、部署、联调和回归工作已完成。当前唯一业务验收阻塞为外部 Active Directory 返回 52e invalidCredentials；正式 HTTPS 还需外部域名与可信证书。"), size=11, color=INK)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        font_run(p.add_run(bold_prefix), bold=True)
        font_run(p.add_run(text[len(bold_prefix):]))
    else:
        font_run(p.add_run(text))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        font_run(p.add_run(item))


def add_table(doc, headers, rows, widths, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_shading(table.cell(0, idx), GRAY)
        p = table.cell(0, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_run(p.add_run(header), size=9.5, bold=True, color=BLUE_DARK)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == status_col else WD_ALIGN_PARAGRAPH.LEFT
            color = INK
            bold = False
            if idx == status_col:
                val = str(value)
                color = GREEN if "PASS" in val else RED if "BLOCK" in val or "FAIL" in val else AMBER
                bold = True
            font_run(p.add_run(str(value)), size=9.2, bold=bold, color=color)
        for idx, cell in enumerate(cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
    return table


def add_figure(doc, filename, caption):
    path = SCREENSHOTS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    with Image.open(path) as image:
        source_width, source_height = image.size
    max_width = 6.35
    max_height = 7.0
    scale = min(max_width / source_width, max_height / source_height)
    run.add_picture(
        str(path),
        width=Inches(source_width * scale),
        height=Inches(source_height * scale),
    )
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(8)
    font_run(cp.add_run(caption), size=9, color=MUTED)


def build():
    REPORTS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    doc.core_properties.title = "海智产品中心 Phase 2 业务实施验收报告"
    doc.core_properties.subject = "AI问答、智能配单、项目BOM、Excel模板导出、资料中心"
    doc.core_properties.author = "海智产品中心项目组"
    doc.core_properties.comments = "正式验收材料"

    add_title_block(doc)
    doc.add_page_break()

    add_heading(doc, "1. 执行摘要", 1)
    add_para(doc, "本轮完成 AI 知识问答、智能配单、项目 BOM、规则校验、不可变版本、Excel 模板映射与导出、资料预览与知识同步、Dify/RAG 集成、权限控制、生产切换和回滚演练。所有可由当前环境独立完成的软件工作均已闭环。")
    add_para(doc, "正式结论为 PHASE 2 BUSINESS IMPLEMENTATION NOT READY。阻塞不是平台缺陷，而是外部 AD 绑定账户认证失败：服务器返回子码 52e（invalidCredentials）。在域管理员确认账户密码、启用、锁定和过期状态前，无法证明销售 LDAP 登录链路。")
    add_table(doc, ["验收门", "结果", "说明"], [
        ("平台技术可用", "PASS", "443 正式服务、健康检查、容器安全、迁移与浏览器回归通过"),
        ("核心业务闭环", "PASS", "问答、配单、BOM、Excel、资料业务闭环通过"),
        ("Dify / LLM / RAG", "PASS", "真实数据集检索、知识同步、模型调用和 80 入口通过"),
        ("安全 / 备份 / 回滚", "PASS", "权限、字段级价格控制、备份与完整回滚恢复演练通过"),
        ("外部目录登录", "BLOCKED", "AD 返回 52e invalidCredentials，需外部凭据修复"),
    ], [2100, 1500, 5760], status_col=1)

    add_heading(doc, "2. 正式运行基线", 1)
    add_table(doc, ["项目", "正式值"], [
        ("海智产品中心", "http://10.1.2.1:443/"),
        ("Dify 正式入口", "http://10.1.2.1/"),
        ("Dify 保留入口", "http://10.1.2.1:18081/"),
        ("Moodle", "http://10.1.2.1:443/moodle/"),
        ("正式镜像", "haizhi-hub-api:6.0.5-phase2-style-fix"),
        ("Git branch", "codex/phase2-business-implementation"),
        ("实现 commit", "1c91a3a70ff8e48da710875e2f7c517634e89d2a"),
        ("Alembic", "c3d4e5f60718 (head)"),
        ("API 回滚容器", "haizhi-hub-api-v604-stylerollback-20260824"),
    ], [2500, 6860])
    add_bullets(doc, [
        "正式 API 以 UID/GID 10001:10001 运行，根文件系统只读。",
        "容器 capability 全部移除，启用 no-new-privileges，/tmp 使用 noexec/nosuid。",
        "443 在开发、候选验收和切换期间保持可用。",
    ])

    add_heading(doc, "3. Phase 2 交付范围", 1)
    add_table(doc, ["模块", "已交付能力", "结果"], [
        ("首页 AI 问答", "结构化关系优先、Dify 资料补充、思考内容清理、来源受控", "PASS"),
        ("智能配单", "自然语言解析、需求确认、真实方案与产品推荐", "PASS"),
        ("项目 BOM", "人工编辑优先、规则校验、不可变版本、价格权限", "PASS"),
        ("Excel", "占位符、可视化单元格映射、BOM 列映射、预览、XLSX", "PASS"),
        ("资料中心", "上传、派生预览、原文件下载、发布、知识同步治理", "PASS"),
        ("Dify / RAG", "单一现有实例、真实 Dataset、80/18081 双入口", "PASS"),
        ("六大中心", "列表与详情回归、数据与权限不回退", "PASS"),
    ], [1700, 5960, 1700], status_col=2)

    doc.add_page_break()
    add_heading(doc, "4. AI 知识问答与 Dify/RAG", 1)
    add_para(doc, "AI Gateway 使用本地模型服务 http://10.1.2.4:6999 和模型 /model/models/Qwen3.6-27B，并对调用身份、模型范围、输出内容和资料来源进行治理。系统先查询六大中心结构化关系，再调用 Dify 当前有效资料；模型输出中的 reasoning_content、<think> 及已知推理泄露标记会被剔除，异常时使用确定性结构化/RAG 结果回退。")
    add_table(doc, ["验证项", "结果", "证据"], [
        ("模型发现与聊天", "PASS", "允许模型约束和真实聊天响应"),
        ("Dify Provider", "PASS", "OpenAI API Compatible 插件 0.0.62"),
        ("Dataset retrieve", "PASS", "Dataset 6b164760-8c9b-4c37-b0f7-96864d88b9c3"),
        ("真实资料同步", "PASS", "文档 14：PUBLISHED / CURRENT / SYNCED"),
        ("草稿同步阻断", "PASS", "未发布资料不可进入知识库"),
        ("Dify 80 入口", "PASS", "setup=finished；原 18081 保留"),
    ], [2700, 1400, 5260], status_col=1)
    add_figure(doc, "01-home-ai-fixed-1920x1080.png", "图 1  首页 AI 知识问答与真实业务信息")

    doc.add_page_break()
    add_heading(doc, "5. 智能配单与需求解析", 1)
    add_para(doc, "验收使用真实桥梁防撞业务语义：某桥梁上下游各 3 公里，需要 4 个球机，支持 AIS 融合、船名 OCR、偏航预警并 7x24 小时运行。解析器对“上下游各3公里”分别写入上游 3 公里、下游 3 公里，不再误合并为单方向距离。")
    add_table(doc, ["流程节点", "输出", "结果"], [
        ("自然语言解析", "距离、设备数量、能力、运行要求结构化", "PASS"),
        ("需求确认", "用户确认后进入方案/BOM 推荐", "PASS"),
        ("真实推荐", "引用现有产品、型号和关联关系", "PASS"),
        ("项目保存", "桥梁防撞业务验收项目-20260824，id 6", "PASS"),
    ], [2200, 5460, 1700], status_col=2)
    add_figure(doc, "02-smart-config-parse-1920x1080.png", "图 2  桥梁防撞需求解析与确认")
    add_figure(doc, "03-recommended-bom-1920x1080.png", "图 3  基于真实关系的推荐 BOM")

    doc.add_page_break()
    add_heading(doc, "6. BOM 引擎、人工编辑与版本", 1)
    add_para(doc, "BOM 编辑器保持人工修改优先。用户修改数量、型号、价格、用途或备注后，记录 manualEdited=true；“重新校验”只运行规则，不覆盖人工内容。主动重新生成推荐 BOM 需要二次确认并创建新版本。")
    add_table(doc, ["验收事实", "结果"], [
        ("当前版本", "V3"),
        ("首条数量", "从 1 人工改为 2"),
        ("人工备注", "人工调整：主平台双机冗余"),
        ("校验统计", "3 PASS / 0 WARNING / 0 ERROR"),
        ("历史保护", "V2、V3 均保留，历史不可覆盖"),
        ("ERROR 覆盖", "销售不可绕过；管理员/产品经理需填写 Override Reason"),
    ], [2900, 6460])
    add_figure(doc, "04-bom-editor-manual-v3-1920x1080.png", "图 4  项目 BOM V3 人工调整与规则校验")

    doc.add_page_break()
    add_heading(doc, "7. Excel 模板映射与导出", 1)
    add_para(doc, "系统既支持带占位符模板，也支持用户现有、无占位符的普通 Excel 模板。普通模板通过多 Sheet 预览、点击单元格映射项目字段，并单独映射 BOM 起始行和各列；最终文件基于上传模板填充，不重新生成“类似格式”的工作簿。")
    add_table(doc, ["能力", "结果", "说明"], [
        ("模板上传/分析", "PASS", "xlsx、多 Sheet、占位符与普通模板"),
        ("项目字段映射", "PASS", "单元格点击映射"),
        ("BOM 列映射", "PASS", "行号、产品、型号、数量、单位、价格、小计、用途、备注"),
        ("数据调整", "PASS", "调整后创建新 BOM 版本"),
        ("效果预览", "PASS", "只读 HTML 预览"),
        ("最终 XLSX", "PASS", "鉴权下载、内容/样式/公式检查"),
    ], [2500, 1400, 5460], status_col=1)
    add_figure(doc, "05-excel-template-mapping-1920x1080.png", "图 5  普通 Excel 模板可视化映射")
    add_figure(doc, "06-excel-effect-preview-1920x1080.png", "图 6  Excel 最终效果预览")
    add_para(doc, "最终样例：export-samples/桥梁防撞业务验收项目-BOM-V3-标准模板.xlsx。样例包含 3 条真实产品、首项数量 2、人工备注、价格合计，并保持合并单元格、冻结窗格、列宽、样式和公式。")

    doc.add_page_break()
    add_heading(doc, "8. 资料中心与知识治理", 1)
    add_para(doc, "原始文件和预览衍生物分开存储。PDF、图片和 TXT 可直接预览；Word/PPT 转换为 PDF；Excel 转换为只读 HTML。在线预览和原文件下载都必须通过登录鉴权，下载返回原始文件而非预览副本。")
    add_table(doc, ["验证项", "结果", "说明"], [
        ("DOCX 预览", "PASS", "返回派生 PDF"),
        ("DOCX 原文件下载", "PASS", "字节保持"),
        ("草稿同步", "PASS", "正确拒绝"),
        ("发布同步", "PASS", "文档 14：PUBLISHED / CURRENT / SYNCED"),
        ("版本治理", "PASS", "CURRENT/HISTORICAL/OBSOLETE 与 OUTDATED"),
        ("脏资料清理", "PASS", "hz-storage-test.txt 已删除"),
    ], [2500, 1400, 5460], status_col=1)
    add_para(doc, "进入 AI 知识库必须同时满足 status=PUBLISHED、documentStatus=CURRENT、knowledgeEnabled=true。文件上传成功本身不会触发知识同步。")
    add_figure(doc, "07-document-center-published-synced-1920x1080.png", "图 7  资料发布、当前版本与知识同步状态")

    doc.add_page_break()
    add_heading(doc, "9. 权限与安全验收", 1)
    add_table(doc, ["角色", "允许", "禁止/约束"], [
        ("普通用户", "知识问答、六大中心只读、已发布资料预览/下载", "无价格、项目、BOM、Excel、后台管理"),
        ("销售", "项目配单、BOM 编辑、授权价格查看", "ERROR 不可绕过"),
        ("产品经理", "知识维护、资料发布同步、项目/BOM/Excel", "敏感字段仍按明确权限"),
        ("管理员", "用户、角色、AD、规则和全部业务能力", "覆盖规则需审计原因"),
    ], [1800, 3920, 3640])
    add_bullets(doc, [
        "后端真实鉴权通过；普通用户写操作返回 403。",
        "无 PRICE_VIEW 时，产品价格和方案 BOM 单价/小计/总价不返回。",
        "文件预览和下载必须登录；Dify 不接收成本价、用户密码和 LDAP 密钥。",
        "管理员后台实现 AD 手工配置、启用、连接测试、候选拉取、选择确认和中文角色分配。",
    ])

    add_heading(doc, "10. 浏览器与响应式回归", 1)
    add_table(doc, ["范围", "结果", "说明"], [
        ("四种视口", "44/44 PASS", "1920x1080、1600x900、1440x900、1366x768"),
        ("六大中心", "12/12 PASS", "6 个列表页 + 6 个详情页"),
        ("横向溢出", "PASS", "首页及六大中心均无异常溢出"),
        ("Console", "PASS", "0 error / 0 warning"),
        ("本地构建", "PASS", "Python compile、frontend build"),
    ], [2500, 1800, 5060], status_col=1)
    add_figure(doc, "final-responsive/home-1366x768.jpg", "图 8  首页 1366x768 最小验收视口")

    doc.add_page_break()
    add_heading(doc, "11. 数据库、备份与回滚", 1)
    add_para(doc, "Phase 2 通过 Alembic revision c3d4e5f60718 新增需求版本、规则、BOM、校验运行、Excel 模板/版本/导出、资料预览和知识同步相关结构。迁移先在隔离数据库验证 upgrade、downgrade 到 7a8c9d0e1f23、再 upgrade 到 head，之后才进入正式环境。")
    add_table(doc, ["项目", "路径/状态"], [
        ("主切换前备份", "/data/haizhi-product-hub/backups/20260824-phase2-final-precutover"),
        ("Dify 80 切换备份", "/data/haizhi-product-hub/backups/20260824-phase2-port80-cutover"),
        ("备份内容", "PostgreSQL、MinIO、运行配置、Dify PostgreSQL、Dify 配置卷"),
        ("校验", "SHA-256 已记录并验证"),
        ("API 回滚", "6.0.4 停止容器保留，可按相同安全参数恢复"),
        ("完整演练", "5.7.3 回滚及 Phase 2 恢复均 PASS"),
    ], [2900, 6460])
    add_para(doc, "旧 port 80 Python/SQLite 服务仅在完整备份后 stop/disable，代码、.env、SQLite 和 systemd unit 均保留。回滚时停止 haizhi-dify-port80，重新 enable/start haizhi-product-center，再验证端口与数据。")

    add_heading(doc, "12. E2E 验收矩阵", 1)
    matrix = [
        ("首页 AI 知识问答", "PASS"), ("AI 智能配单", "PASS"),
        ("需求确认", "PASS"), ("真实 BOM 推荐", "PASS"),
        ("BOM 规则", "PASS"), ("人工 BOM 编辑", "PASS"),
        ("人工修改不被覆盖", "PASS"), ("BOM 版本", "PASS"),
        ("项目配单", "PASS"), ("价格权限", "PASS"),
        ("Excel 模板上传", "PASS"), ("Excel 可视化映射", "PASS"),
        ("BOM 列映射", "PASS"), ("Excel 数据调整", "PASS"),
        ("Excel 效果预览", "PASS"), ("最终 XLSX 导出/格式", "PASS"),
        ("资料上传/预览/下载", "PASS"), ("资料知识同步规则", "PASS"),
        ("Dify API / 80 入口", "PASS"), ("六大中心保持正常", "PASS"),
        ("普通/销售/产品经理/管理员权限", "PASS"), ("备份恢复/回滚方案", "PASS"),
        ("浏览器回归", "PASS"), ("销售 LDAP 实际登录", "BLOCKED"),
    ]
    add_table(doc, ["验收项", "结果"], matrix, [7160, 2200], status_col=1)

    doc.add_page_break()
    add_heading(doc, "13. 外部阻塞与已知限制", 1)
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360])
    set_cell_shading(callout.cell(0, 0), "FDECEC")
    p = callout.cell(0, 0).paragraphs[0]
    font_run(p.add_run("唯一业务验收阻塞："), size=11, bold=True, color=RED)
    font_run(p.add_run("Active Directory 10.1.1.102 的 LDAP bind 返回 LDAPInvalidCredentialsResult / data 52e。网络、389/636、RootDSE、域名和 Base DN 已验证，平台无法在没有有效域凭据的情况下完成销售 LDAP 登录。"), size=11)
    add_para(doc, "外部处理要求：由域管理员确认 ldapreader@hilaicloud.com 的准确密码、账户是否启用、是否锁定、密码是否过期以及登录限制。修复后在系统管理中重新执行“连接测试”，再手动同步候选用户、选择销售用户并确认授权，最后完成销售 LDAP 登录 E2E。")
    add_para(doc, "HTTPS 限制：当前地址使用 http://10.1.2.1:443/。可信 HTTPS 需要外部提供正式域名、DNS 解析和可用证书。该项不影响当前 HTTP 业务验证，但属于正式对外安全发布前置条件。")
    add_para(doc, "Dify 控制台密码登录采用客户端公钥加密，简单 HTTP 明文密码探测会返回 Invalid encrypted data；现有 Provider、应用调用、资料同步和检索均已通过，不属于运行阻塞。")
    add_para(doc, "账号清理限制：临时浏览器验收账号 id 15 存在审计历史，物理 DELETE 会被数据库审计外键阻止。该账号已通过管理员接口禁用并复核；生产运维应对有审计记录的账号采用禁用保留，而不是物理删除。")

    add_heading(doc, "14. 最终结论", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    font_run(p.add_run("PHASE 2 BUSINESS IMPLEMENTATION NOT READY"), size=16, bold=True, color=RED)
    add_para(doc, "除 AD 实际认证与正式域名/证书外，本轮范围内的开发、部署、业务闭环、数据治理、Dify/RAG、权限、浏览器、备份和回滚验证均已完成。AD 52e 修复并补测销售 LDAP 登录后，可重新判定是否达到 READY。")

    add_heading(doc, "附录 A  交付物索引", 1)
    add_bullets(doc, [
        "Markdown 报告：reports/ 目录下 9 份专项报告。",
        "截图证据：screenshots/，含核心业务、四视口和六大中心 12 条路由。",
        "Excel 样例：export-samples/桥梁防撞业务验收项目-BOM-V3-标准模板.xlsx。",
        "测试模板与检查记录：tests/；恢复事实源：项目根目录 CODEX_CHECKPOINT.md。",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
