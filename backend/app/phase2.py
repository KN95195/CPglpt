import hashlib
import io
import json
import os
import re
import urllib.parse
import uuid
from copy import copy
from datetime import datetime, timezone
from typing import Any

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import settings
from .database import session
from .dify_knowledge import DifyKnowledgeClient
from .document_preview import preview_capable,xlsx_preview_html
from .models import *
from .storage import storage

PROJECT_STATUSES={'DRAFT','CONFIGURING','CONFIRMED','EXPORTED','ARCHIVED'}
RULE_TYPES={'REQUIRE','QUANTITY','CONDITION','COMPATIBILITY','EXCLUDE','RECOMMEND'}
BOM_COLUMNS={'lineNo','productName','model','quantity','unit','unitPrice','subtotal','purpose','note'}

_PRIVATE_REASONING_MARKERS=(
    "here's a thinking process",
    'analyze user input',
    'internal refinement',
    'final output generation',
    'self-correction/verification',
)

def public_ai_answer(content:str|None,fallback:str)->str:
    """Return only user-facing model content and fail closed on reasoning leakage."""
    text=(content or '').strip()
    if not text:return fallback
    text=re.sub(r'<think>.*?</think>','',text,flags=re.IGNORECASE|re.DOTALL).strip()
    if '</think>' in text.lower():text=re.split(r'</think>',text,flags=re.IGNORECASE)[-1].strip()
    final_match=re.search(r'(?:final answer|最终答案)\s*[:：]\s*(.+)$',text,flags=re.IGNORECASE|re.DOTALL)
    if final_match:text=final_match.group(1).strip()
    lowered=text.lower()
    if not text or any(marker in lowered for marker in _PRIVATE_REASONING_MARKERS):return fallback
    return text

def parse_requirement_text(raw:str)->dict[str,Any]:
    def number(pattern,default=0):
        match=re.search(pattern,raw);return int(match.group(1)) if match else default
    shared_distance=number(r'上下游\D{0,4}各\D{0,4}(\d+)')
    return {
        'scene':'桥梁防撞' if '桥' in raw else '待确认',
        'upstreamKm':number(r'上游\D{0,4}(\d+)',shared_distance),
        'downstreamKm':number(r'下游\D{0,4}(\d+)',shared_distance),
        'ptzCount':number(r'(\d+)\s*个?球机'),
        'ais':'AIS' in raw.upper(),
        'ocr':'OCR' in raw.upper() or '船名' in raw,
        'yaw':'偏航' in raw,
        'continuousOperation':'7×24' in raw or '7x24' in raw.lower(),
    }

class AiChatIn(BaseModel):
    message:str=Field(min_length=2,max_length=4000)
    conversation_id:str|None=None
class RequirementParseIn(BaseModel):
    raw_requirement:str=Field(min_length=5,max_length=10000)
class RequirementConfirmIn(BaseModel):
    project_id:int
    raw_requirement:str
    ai_parsed_requirement:dict[str,Any]={}
    user_adjusted_requirement:dict[str,Any]={}
    confirmed_requirement:dict[str,Any]
class ProjectCreate(BaseModel):
    name:str=Field(min_length=2,max_length=160);customer:str=Field(min_length=1,max_length=160);region:str='';scene_id:int;notes:str='';raw_requirement:str='';requirements:dict[str,Any]={}
class ProjectUpdate(BaseModel):
    name:str|None=None;customer:str|None=None;region:str|None=None;scene_id:int|None=None;status:str|None=None;notes:str|None=None
class BomRecommendIn(BaseModel):
    project_id:int;requirement_version_id:int|None=None
class BomSaveIn(BaseModel):
    items:list[dict[str,Any]];change_summary:str='人工调整';source:str='MANUAL_EDIT'
class BomValidateIn(BaseModel):
    project_id:int;bom_version_id:int|None=None;override_reason:str=''
class RuleIn(BaseModel):
    code:str=Field(pattern=r'^[A-Z0-9_]{2,80}$');name:str=Field(min_length=2,max_length=160);rule_type:str;condition:dict[str,Any];action:dict[str,Any];enabled:bool=True
class RulePatch(BaseModel):
    name:str|None=None;rule_type:str|None=None;condition:dict[str,Any]|None=None;action:dict[str,Any]|None=None;enabled:bool|None=None
class TemplateMappingIn(BaseModel):
    sheet_name:str;placeholder_mappings:dict[str,str]={};bom_template_row:int=Field(ge=1);bom_column_mappings:dict[str,str]
class DocumentPatch(BaseModel):
    name:str|None=None;category:str|None=None;version:str|None=None;description:str|None=None;status:str|None=None;document_status:str|None=None;applicable_models:str|None=None;applicable_versions:str|None=None;knowledge_enabled:bool|None=None;center_type:str|None=None;center_id:int|None=None

def build_phase2_router(current_user, permit, permission_codes, audit):
    router=APIRouter()

    def json_value(value,default):
        try:return json.loads(value or '')
        except Exception:return default
    def require_project_access(project,u):
        codes=permission_codes(u)
        if 'BOM_VIEW' not in codes:raise HTTPException(403,'无项目配单权限')
        if u.role.code=='sales' and project.sales_owner_id not in (None,u.id):raise HTTPException(403,'只能访问本人负责的项目')
    def product_line(product,quantity=1,purpose='',reason='',unit_price=None):
        return {'productId':product.id,'productName':product.name,'model':product.model_code,'quantity':quantity,'unit':'台','unitPrice':unit_price,'referencePrice':None,'purpose':purpose,'note':reason,'manualEdited':False}
    def price_for(db,product_id,u):
        if 'PRICE_VIEW' not in permission_codes(u):return None
        price=db.scalar(select(ProductPrice).where(ProductPrice.product_id==product_id))
        return float(price.reference_price) if price else None
    def bom_dto(version,u):
        items=json_value(version.snapshot_json,[])
        if 'PRICE_VIEW' not in permission_codes(u):
            for item in items:item.pop('unitPrice',None);item.pop('referencePrice',None);item.pop('subtotal',None)
        return {'id':version.id,'version':version.version,'requirementVersionId':version.requirement_version_id,'source':version.source,'items':items,'manualEdited':version.manual_edited,'changeSummary':version.change_summary,'createdAt':version.created_at.isoformat()}
    def document_dto(rec,u):
        return {'id':rec.id,'name':rec.name,'originalFileName':rec.original_file_name or rec.name,'mimeType':rec.mime_type,'fileSize':rec.file_size,'category':rec.category,'version':rec.version,'description':rec.description,'status':rec.status,'documentStatus':rec.document_status,'applicableModels':rec.applicable_models,'applicableVersions':rec.applicable_versions,'knowledgeEnabled':rec.knowledge_enabled,'knowledgeStatus':rec.knowledge_status,'knowledgeSyncError':rec.knowledge_sync_error,'centerType':rec.center_type,'centerId':rec.center_id,'canPreview':preview_capable(rec.original_file_name or rec.name,rec.mime_type),'canDownload':'DOCUMENT_DOWNLOAD' in permission_codes(u),'updatedAt':rec.updated_at.isoformat()}

    @router.post('/api/ai/chat')
    async def ai_chat(x:AiChatIn,db:Session=Depends(session),u=Depends(current_user)):
        query=x.message.strip();terms=[term for term in re.split(r'[\s，。？、,]+',query) if len(term)>=2]
        products=list(db.scalars(select(Product).order_by(Product.updated_at.desc())))
        matched=[p for p in products if any(term in (p.name+p.model_code+p.summary) for term in terms)]
        if not matched and ('终端' in query or '算法' in query):matched=[p for p in products if '终端' in p.name][:3]
        sources=[];facts=[]
        for p in matched[:5]:
            relations=[]
            for rel in db.scalars(select(KnowledgeRelation).where(((KnowledgeRelation.source_type=='products')&(KnowledgeRelation.source_id==p.id))|((KnowledgeRelation.target_type=='products')&(KnowledgeRelation.target_id==p.id)))):
                other_type=rel.target_type if rel.source_type=='products' else rel.source_type;other_id=rel.target_id if rel.source_type=='products' else rel.source_id
                model={'algorithms':Algorithm,'software':Software,'model-capabilities':Capability,'scenes':Scene,'solutions':Solution}.get(other_type);obj=db.get(model,other_id) if model else None
                if obj:relations.append({'type':other_type,'id':obj.id,'name':obj.name})
            facts.append({'id':p.id,'name':p.name,'model':p.model_code,'summary':p.summary,'relations':relations,'referencePrice':price_for(db,p.id,u)})
            sources.append({'type':'products','id':p.id,'name':p.name,'url':f'/products/{p.id}'})
        docs=list(db.scalars(select(DocumentAsset).where(DocumentAsset.status=='PUBLISHED',DocumentAsset.document_status=='CURRENT').order_by(DocumentAsset.updated_at.desc()).limit(50)))
        related_docs=[d for d in docs if any(term in (d.name+d.description+d.applicable_models) for term in terms)][:5]
        rag_excerpts=[]
        try:
            retrieved=await DifyKnowledgeClient().retrieve(query,top_k=5)
            synced_docs={d.knowledge_document_id:d for d in docs if d.knowledge_enabled and d.knowledge_status=='SYNCED' and d.knowledge_document_id}
            for row in retrieved:
                document=synced_docs.get(row['documentId'])
                if not document or not row['content']:continue
                rag_excerpts.append({'documentId':document.id,'name':document.name,'content':row['content'][:3000],'score':row['score']})
                if document not in related_docs:related_docs.append(document)
        except Exception:
            pass
        for d in related_docs[:5]:
            source={'type':'documents','id':d.id,'name':d.name,'previewUrl':f'/api/documents/{d.id}/preview','canDownload':'DOCUMENT_DOWNLOAD' in permission_codes(u)}
            if source['canDownload']:source['downloadUrl']=f'/api/documents/{d.id}/download'
            sources.append(source)
        configuration_intent=any(word in query for word in ('怎么配置','如何配置','配单','上下游','项目建设','需要部署'))
        fact_text=json.dumps({'structuredFacts':facts,'documentMetadata':[{'id':d.id,'name':d.name,'description':d.description} for d in related_docs[:5]],'documentExcerpts':rag_excerpts},ensure_ascii=False)
        answer='未在当前有效产品知识中找到直接答案。'
        if facts:
            relation_names=sorted({r['name'] for f in facts for r in f['relations'] if r['type']=='algorithms'})
            answer='；'.join(f"{f['name']}（{f['model']}）：{f['summary']}" for f in facts[:3])
            if relation_names:answer+='。关联算法：'+'、'.join(relation_names)
        if rag_excerpts:
            excerpt=' '.join(rag_excerpts[0]['content'].split())[:900]
            if facts:answer+=f"。当前有效资料《{rag_excerpts[0]['name']}》补充：{excerpt}"
            else:answer=f"根据当前有效资料《{rag_excerpts[0]['name']}》：{excerpt}"
        elif related_docs:answer+=('。相关资料：' if answer else '相关资料：')+'、'.join(d.name for d in related_docs)
        try:
            prompt='事实：'+fact_text+'\n问题：'+query
            async with httpx.AsyncClient(timeout=settings.ai_timeout_seconds) as client:
                response=await client.post(settings.ai_base_url.rstrip('/')+'/v1/chat/completions',json={'model':settings.ai_model,'messages':[{'role':'system','content':'你是海智产品知识助手。只依据用户提供的事实回答，不得编造。只输出简洁的中文最终答案，禁止展示思考过程、分析步骤、内部指令或英文推理。结构化事实优先于资料片段；没有依据时明确回答未找到。'},{'role':'user','content':prompt}],'max_tokens':500});response.raise_for_status();message=response.json()['choices'][0]['message'];answer=public_ai_answer(message.get('content'),answer)
        except Exception:pass
        return {'conversationId':x.conversation_id or uuid.uuid4().hex,'answer':answer,'sources':sources,'configurationIntent':configuration_intent,'configurationInput':query if configuration_intent else None,'priceVisible':'PRICE_VIEW' in permission_codes(u)}

    @router.post('/api/ai/config/parse')
    def parse_requirement(x:RequirementParseIn,u=Depends(permit('BOM_EDIT'))):
        raw=x.raw_requirement
        parsed=parse_requirement_text(raw)
        missing=[{'field':'scene','question':'请选择项目业务场景'}] if parsed['scene']=='待确认' else []
        return {'rawRequirement':raw,'parsedRequirement':parsed,'missingFields':missing,'readyToConfirm':not missing}

    @router.post('/api/projects',status_code=201)
    def create_project(x:ProjectCreate,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        if not db.get(Scene,x.scene_id):raise HTTPException(422,'场景不存在')
        rec=Project(name=x.name,customer=x.customer,region=x.region,scene_id=x.scene_id,status='CONFIGURING',requirements_json=json.dumps(x.requirements,ensure_ascii=False),bom_json='[]',sales_owner_id=u.id,created_by=u.id,notes=x.notes);db.add(rec);db.flush()
        if x.raw_requirement or x.requirements:
            requirement=ProjectRequirementVersion(project_id=rec.id,version=1,raw_requirement=x.raw_requirement,ai_parsed_json=json.dumps(x.requirements,ensure_ascii=False),user_adjusted_json=json.dumps(x.requirements,ensure_ascii=False),confirmed_json='{}');db.add(requirement);db.flush();rec.current_requirement_version_id=requirement.id
        audit(db,u,'CREATE','project',rec.id,{'status':rec.status});db.commit();db.refresh(rec);return {'id':rec.id,'name':rec.name,'status':rec.status}

    @router.post('/api/ai/config/confirm')
    def confirm_requirement(x:RequirementConfirmIn,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,x.project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);version=(db.scalar(select(func.max(ProjectRequirementVersion.version)).where(ProjectRequirementVersion.project_id==project.id)) or 0)+1
        rec=ProjectRequirementVersion(project_id=project.id,version=version,raw_requirement=x.raw_requirement,ai_parsed_json=json.dumps(x.ai_parsed_requirement,ensure_ascii=False),user_adjusted_json=json.dumps(x.user_adjusted_requirement,ensure_ascii=False),confirmed_json=json.dumps(x.confirmed_requirement,ensure_ascii=False),confirmed_by=u.id,confirmed_at=datetime.now(timezone.utc));db.add(rec);db.flush();project.current_requirement_version_id=rec.id;project.requirements_json=rec.confirmed_json;project.status='CONFIGURING';audit(db,u,'CONFIRM','project_requirement',rec.id,{'version':version});db.commit();return {'projectId':project.id,'requirementVersionId':rec.id,'version':version,'status':'CONFIRMED'}

    @router.get('/api/projects')
    def list_projects(db:Session=Depends(session),u=Depends(permit('BOM_VIEW'))):
        query=select(Project).order_by(Project.updated_at.desc())
        if u.role.code=='sales':query=query.where(Project.sales_owner_id==u.id)
        return [{'id':p.id,'name':p.name,'customer':p.customer,'region':p.region,'scene':p.scene.name,'status':p.status,'bomVersion':p.bom_version,'updatedAt':p.updated_at.isoformat()} for p in db.scalars(query)]

    @router.get('/api/projects/{project_id}')
    def project_detail(project_id:int,db:Session=Depends(session),u=Depends(permit('BOM_VIEW'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);version=db.get(ProjectBomVersion,project.current_bom_version_id) if project.current_bom_version_id else None
        return {'id':project.id,'name':project.name,'customer':project.customer,'region':project.region,'sceneId':project.scene_id,'scene':project.scene.name,'status':project.status,'notes':project.notes,'currentRequirementVersionId':project.current_requirement_version_id,'currentBomVersion':bom_dto(version,u) if version else None,'updatedAt':project.updated_at.isoformat()}

    @router.patch('/api/projects/{project_id}')
    def patch_project(project_id:int,x:ProjectUpdate,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);values=x.model_dump(exclude_none=True)
        if 'status' in values and values['status'] not in PROJECT_STATUSES:raise HTTPException(422,'项目状态无效')
        if 'scene_id' in values and not db.get(Scene,values['scene_id']):raise HTTPException(422,'场景不存在')
        for key,value in values.items():setattr(project,key,value)
        audit(db,u,'UPDATE','project',project.id,values);db.commit();return {'id':project.id,'status':project.status}

    @router.post('/api/bom/recommend')
    def recommend_bom(x:BomRecommendIn,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,x.project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);requirement=db.get(ProjectRequirementVersion,x.requirement_version_id or project.current_requirement_version_id)
        if not requirement or requirement.project_id!=project.id:raise HTTPException(422,'请先确认项目需求')
        solution=db.scalar(select(Solution).where(Solution.scene_id==project.scene_id,Solution.status=='ACTIVE').order_by(Solution.id))
        solution_items=list(db.scalars(select(SolutionBomItem).where(SolutionBomItem.solution_id==solution.id).order_by(SolutionBomItem.id))) if solution else []
        items=[]
        for row in solution_items:
            price=price_for(db,row.product_id,u);item=product_line(row.product,row.quantity,row.purpose,row.recommendation_reason,price);item['referencePrice']=price;items.append(item)
        if not items:raise HTTPException(422,'当前场景没有可用标准方案BOM，请由产品经理维护方案BOM后重试')
        number=(db.scalar(select(func.max(ProjectBomVersion.version)).where(ProjectBomVersion.project_id==project.id)) or 0)+1
        version=ProjectBomVersion(project_id=project.id,version=number,requirement_version_id=requirement.id,source='AI_RECOMMEND',snapshot_json=json.dumps(items,ensure_ascii=False),manual_edited=False,change_summary='基于确认需求和标准方案生成',created_by=u.id);db.add(version);db.flush();project.current_bom_version_id=version.id;project.bom_version=number;project.bom_json=version.snapshot_json;project.status='CONFIGURING';audit(db,u,'CREATE','project_bom_version',version.id,{'version':number,'solutionId':solution.id if solution else None});db.commit();return bom_dto(version,u)

    @router.post('/api/projects/{project_id}/bom-versions',status_code=201)
    def save_bom_version(project_id:int,x:BomSaveIn,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u)
        for item in x.items:
            product=db.get(Product,item.get('productId'))
            if not product:raise HTTPException(422,'BOM只能引用真实产品记录')
            if item.get('model') not in {product.model_code,*[v.model_code for v in product.variants]}:raise HTTPException(422,f'{product.name}的型号不属于当前产品')
            item['productName']=product.name;item['manualEdited']=True
        number=(db.scalar(select(func.max(ProjectBomVersion.version)).where(ProjectBomVersion.project_id==project.id)) or 0)+1
        version=ProjectBomVersion(project_id=project.id,version=number,requirement_version_id=project.current_requirement_version_id,source=x.source,snapshot_json=json.dumps(x.items,ensure_ascii=False),manual_edited=True,change_summary=x.change_summary,created_by=u.id);db.add(version);db.flush();project.current_bom_version_id=version.id;project.bom_version=number;project.bom_json=version.snapshot_json;audit(db,u,'CREATE','project_bom_version',version.id,{'manualEdited':True});db.commit();return bom_dto(version,u)

    @router.get('/api/projects/{project_id}/bom-versions')
    def bom_versions(project_id:int,db:Session=Depends(session),u=Depends(permit('BOM_VIEW'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);return [bom_dto(v,u) for v in db.scalars(select(ProjectBomVersion).where(ProjectBomVersion.project_id==project_id).order_by(ProjectBomVersion.version.desc()))]

    @router.post('/api/bom/validate')
    def validate_bom(x:BomValidateIn,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,x.project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);version=db.get(ProjectBomVersion,x.bom_version_id or project.current_bom_version_id)
        if not version or version.project_id!=project.id:raise HTTPException(422,'BOM版本不存在')
        items=json_value(version.snapshot_json,[]);results=[]
        for item in items:
            level='PASS';message='产品和型号有效'
            product=db.get(Product,item.get('productId'))
            if not product:level='ERROR';message='产品记录不存在'
            elif int(item.get('quantity') or 0)<=0:level='ERROR';message='数量必须大于0'
            results.append({'level':level,'productId':item.get('productId'),'message':message})
        for rule in db.scalars(select(BomRule).where(BomRule.enabled==True)):
            condition=json_value(rule.condition_json,{});action=json_value(rule.recommendation_json,{})
            if condition.get('sceneId') not in (None,project.scene_id):continue
            product_id=action.get('productId');present=any(item.get('productId')==product_id for item in items)
            if rule.rule_type=='REQUIRE' and product_id and not present:results.append({'level':'ERROR','ruleId':rule.id,'message':f'必须包含：{db.get(Product,product_id).name if db.get(Product,product_id) else product_id}'})
            if rule.rule_type=='RECOMMEND' and product_id and not present:results.append({'level':'WARNING','ruleId':rule.id,'message':'建议补充规则指定产品'})
        errors=sum(r['level']=='ERROR' for r in results);warnings=sum(r['level']=='WARNING' for r in results);passes=sum(r['level']=='PASS' for r in results)
        override_by=None;override_at=None
        if errors and x.override_reason:
            if u.role.code not in {'admin','product_admin'}:raise HTTPException(403,'销售不能绕过错误校验')
            override_by=u.id;override_at=datetime.now(timezone.utc)
        run=BomValidationRun(project_id=project.id,bom_version_id=version.id,run_by=u.id,rule_version_set_json=json.dumps([{'id':r.id,'version':r.current_version} for r in db.scalars(select(BomRule).where(BomRule.enabled==True))]),pass_count=passes,warning_count=warnings,error_count=errors,result_json=json.dumps(results,ensure_ascii=False),override_by=override_by,override_reason=x.override_reason if override_by else '',override_at=override_at);db.add(run);audit(db,u,'VALIDATE','project_bom_version',version.id,{'errors':errors,'warnings':warnings,'overridden':bool(override_by)});db.commit();return {'id':run.id,'status':'ERROR' if errors and not override_by else 'WARNING' if warnings else 'PASS','passCount':passes,'warningCount':warnings,'errorCount':errors,'results':results,'overridden':bool(override_by)}

    @router.get('/api/bom/rules')
    def list_rules(db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        return [{'id':r.id,'code':r.code,'name':r.name,'ruleType':r.rule_type,'condition':json_value(r.condition_json,{}),'action':json_value(r.recommendation_json,{}),'enabled':r.enabled,'version':r.current_version} for r in db.scalars(select(BomRule).order_by(BomRule.code))]

    @router.post('/api/bom/rules',status_code=201)
    def create_rule(x:RuleIn,db:Session=Depends(session),u=Depends(permit('KNOWLEDGE_MANAGE'))):
        if x.rule_type not in RULE_TYPES:raise HTTPException(422,'规则类型无效')
        if db.scalar(select(BomRule).where(BomRule.code==x.code)):raise HTTPException(409,'规则编码已存在')
        rule=BomRule(code=x.code,name=x.name,rule_type=x.rule_type,condition_json=json.dumps(x.condition,ensure_ascii=False),recommendation_json=json.dumps(x.action,ensure_ascii=False),enabled=x.enabled,current_version=1);db.add(rule);db.flush();db.add(BomRuleVersion(rule_id=rule.id,version=1,rule_type=x.rule_type,condition_json=rule.condition_json,action_json=rule.recommendation_json,enabled=x.enabled,changed_by=u.id));audit(db,u,'CREATE','bom_rule',rule.id,{'version':1});db.commit();return {'id':rule.id,'version':1}

    @router.patch('/api/bom/rules/{rule_id}')
    def patch_rule(rule_id:int,x:RulePatch,db:Session=Depends(session),u=Depends(permit('KNOWLEDGE_MANAGE'))):
        rule=db.get(BomRule,rule_id)
        if not rule:raise HTTPException(404,'规则不存在')
        values=x.model_dump(exclude_none=True)
        if values.get('rule_type') and values['rule_type'] not in RULE_TYPES:raise HTTPException(422,'规则类型无效')
        if 'condition' in values:rule.condition_json=json.dumps(values.pop('condition'),ensure_ascii=False)
        if 'action' in values:rule.recommendation_json=json.dumps(values.pop('action'),ensure_ascii=False)
        for key,value in values.items():setattr(rule,key,value)
        rule.current_version+=1;db.add(BomRuleVersion(rule_id=rule.id,version=rule.current_version,rule_type=rule.rule_type,condition_json=rule.condition_json,action_json=rule.recommendation_json,enabled=rule.enabled,changed_by=u.id));audit(db,u,'UPDATE','bom_rule',rule.id,{'version':rule.current_version});db.commit();return {'id':rule.id,'version':rule.current_version}

    @router.get('/api/excel-templates')
    def list_templates(db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        return [{'id':t.id,'name':t.name,'description':t.description,'currentVersion':t.current_version,'enabled':t.enabled,'updatedAt':t.updated_at.isoformat()} for t in db.scalars(select(ExcelTemplate).order_by(ExcelTemplate.updated_at.desc()))]

    @router.post('/api/excel-templates',status_code=201)
    async def upload_template(file:UploadFile=File(...),name:str=Form(''),description:str=Form(''),db:Session=Depends(session),u=Depends(permit('KNOWLEDGE_MANAGE'))):
        data=await file.read()
        if not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(415,'仅支持.xlsx模板')
        try:workbook=load_workbook(io.BytesIO(data),data_only=False)
        except Exception as exc:raise HTTPException(422,'Excel模板无法读取') from exc
        digest=hashlib.sha256(data).hexdigest();object_name=f'excel-templates/{uuid.uuid4().hex}-{os.path.basename(file.filename or "template.xlsx")}';path=storage.put(object_name,data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        template=ExcelTemplate(name=name or os.path.splitext(file.filename or '配单模板')[0],description=description,created_by=u.id);db.add(template);db.flush();version=ExcelTemplateVersion(template_id=template.id,version=1,original_file_name=file.filename or 'template.xlsx',file_path=path,file_hash=digest,sheet_name=workbook.sheetnames[0],uploaded_by=u.id);db.add(version);audit(db,u,'CREATE','excel_template',template.id,{'fileHash':digest});db.commit();return {'id':template.id,'versionId':version.id,'sheetNames':workbook.sheetnames,'fileHash':digest}

    def load_template_bytes(version):
        response=storage.get(version.file_path.split('/',1)[1]);data=response.read();response.close();response.release_conn();return data

    @router.post('/api/excel-templates/{template_id}/analyze')
    def analyze_template(template_id:int,db:Session=Depends(session),u=Depends(permit('KNOWLEDGE_MANAGE'))):
        template=db.get(ExcelTemplate,template_id)
        if not template:raise HTTPException(404,'模板不存在')
        version=db.scalar(select(ExcelTemplateVersion).where(ExcelTemplateVersion.template_id==template.id,ExcelTemplateVersion.version==template.current_version));workbook=load_workbook(io.BytesIO(load_template_bytes(version)),data_only=False)
        placeholders={};preview=[];sheets=[]
        for sheet in workbook.worksheets:
            sheet_cells=[]
            for row in sheet.iter_rows():
                values=[]
                for cell in row:
                    text=str(cell.value or '');values.append(text)
                    for key in re.findall(r'\{\{([a-zA-Z0-9_]+)\}\}',text):placeholders[key]=f'{sheet.title}!{cell.coordinate}'
                    if text and cell.row<=120 and cell.column<=40:sheet_cells.append({'coordinate':cell.coordinate,'row':cell.row,'column':cell.column,'value':text[:300],'bold':bool(cell.font.bold),'align':cell.alignment.horizontal or 'left'})
                if any(values):preview.append({'sheet':sheet.title,'row':row[0].row,'values':values[:20]})
            sheets.append({'name':sheet.title,'maxRow':min(max(sheet.max_row,1),120),'maxColumn':min(max(sheet.max_column,1),40),'cells':sheet_cells,'mergedRanges':[str(item) for item in sheet.merged_cells.ranges]})
        saved_placeholders=json_value(version.placeholder_mappings_json,{})
        result={'sheetNames':workbook.sheetnames,'placeholders':placeholders,'previewRows':preview[:100],'sheets':sheets,'mapping':{'sheetName':version.sheet_name or workbook.sheetnames[0],'placeholderMappings':saved_placeholders or placeholders,'bomTemplateRow':version.bom_template_row or 1,'bomColumnMappings':json_value(version.bom_column_mappings_json,{})}}
        version.analyzed_json=json.dumps(result,ensure_ascii=False)
        if not saved_placeholders:version.placeholder_mappings_json=json.dumps(placeholders,ensure_ascii=False)
        db.commit();return result

    @router.post('/api/excel-templates/{template_id}/mapping')
    def save_mapping(template_id:int,x:TemplateMappingIn,db:Session=Depends(session),u=Depends(permit('KNOWLEDGE_MANAGE'))):
        if set(x.bom_column_mappings)-BOM_COLUMNS:raise HTTPException(422,'包含不支持的BOM字段映射')
        template=db.get(ExcelTemplate,template_id)
        if not template:raise HTTPException(404,'模板不存在')
        version=db.scalar(select(ExcelTemplateVersion).where(ExcelTemplateVersion.template_id==template.id,ExcelTemplateVersion.version==template.current_version));workbook=load_workbook(io.BytesIO(load_template_bytes(version)),data_only=False)
        if x.sheet_name not in workbook.sheetnames:raise HTTPException(422,'BOM工作表不存在')
        allowed_fields={'project_name','customer_name','sales_name','project_region','export_date','total_price'}
        if set(x.placeholder_mappings)-allowed_fields:raise HTTPException(422,'包含不支持的项目字段映射')
        coordinate_pattern=re.compile(r'^(?:(.+)!)?([A-Z]{1,3}[1-9][0-9]*)$',re.IGNORECASE)
        for coordinate in x.placeholder_mappings.values():
            match=coordinate_pattern.fullmatch(coordinate)
            if not match or (match.group(1) and match.group(1) not in workbook.sheetnames):raise HTTPException(422,f'单元格坐标无效：{coordinate}')
        for coordinate in x.bom_column_mappings.values():
            if not re.fullmatch(r'[A-Z]{1,3}|[A-Z]{1,3}[1-9][0-9]*',coordinate.upper()):raise HTTPException(422,f'BOM列坐标无效：{coordinate}')
        version.sheet_name=x.sheet_name;version.placeholder_mappings_json=json.dumps(x.placeholder_mappings,ensure_ascii=False);version.bom_template_row=x.bom_template_row;version.bom_column_mappings_json=json.dumps({key:value.upper() for key,value in x.bom_column_mappings.items()},ensure_ascii=False);audit(db,u,'UPDATE','excel_template_mapping',version.id,x.model_dump());db.commit();return {'templateVersionId':version.id,'saved':True}

    def render_workbook(project,bom_version,template_version):
        workbook=load_workbook(io.BytesIO(load_template_bytes(template_version)),data_only=False);sheet=workbook[template_version.sheet_name or workbook.sheetnames[0]];items=json_value(bom_version.snapshot_json,[])
        fields={'project_name':project.name,'customer_name':project.customer,'sales_name':'','project_region':project.region,'export_date':datetime.now().strftime('%Y-%m-%d'),'total_price':sum(float(i.get('unitPrice') or 0)*int(i.get('quantity') or 0) for i in items)}
        for row in workbook.worksheets:
            for cells in row.iter_rows():
                for cell in cells:
                    if isinstance(cell.value,str):
                        for key,value in fields.items():cell.value=cell.value.replace('{{'+key+'}}',str(value))
        for key,coordinate in json_value(template_version.placeholder_mappings_json,{}).items():
            if key not in fields:continue
            if '!' in coordinate:sheet_name,cell_coordinate=coordinate.rsplit('!',1)
            else:sheet_name,cell_coordinate=sheet.title,coordinate
            if sheet_name in workbook.sheetnames:workbook[sheet_name][cell_coordinate]=fields[key]
        start=template_version.bom_template_row or 1;mappings=json_value(template_version.bom_column_mappings_json,{})
        if len(items)>1:sheet.insert_rows(start+1,len(items)-1)
        for index,item in enumerate(items):
            row_no=start+index
            if index and start>0:
                source=sheet[start]
                for source_cell,target_cell in zip(source,sheet[row_no]):
                    if source_cell.has_style:target_cell._style=copy(source_cell._style)
                    if source_cell.number_format:target_cell.number_format=source_cell.number_format
            values={'lineNo':index+1,'productName':item.get('productName'),'model':item.get('model'),'quantity':item.get('quantity'),'unit':item.get('unit'),'unitPrice':item.get('unitPrice'),'subtotal':float(item.get('unitPrice') or 0)*int(item.get('quantity') or 0),'purpose':item.get('purpose'),'note':item.get('note')}
            for field,column in mappings.items():sheet[f'{column}{row_no}' if re.fullmatch(r'[A-Za-z]+',column) else column.replace(str(start),str(row_no))]=values.get(field)
        output=io.BytesIO();workbook.save(output);return output.getvalue(),fields,items

    @router.post('/api/projects/{project_id}/export-preview')
    def export_preview(project_id:int,template_id:int,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);bom=db.get(ProjectBomVersion,project.current_bom_version_id);template=db.get(ExcelTemplate,template_id)
        if not bom or not template:raise HTTPException(422,'请选择有效BOM和Excel模板')
        version=db.scalar(select(ExcelTemplateVersion).where(ExcelTemplateVersion.template_id==template.id,ExcelTemplateVersion.version==template.current_version));rendered,fields,items=render_workbook(project,bom,version);html=xlsx_preview_html(rendered).decode('utf-8')
        return {'project':fields,'items':items,'template':{'id':template.id,'name':template.name,'version':version.version,'sheetName':version.sheet_name},'html':html,'readOnly':True}

    @router.post('/api/projects/{project_id}/export-excel')
    def export_excel(project_id:int,template_id:int,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        project=db.get(Project,project_id)
        if not project:raise HTTPException(404,'项目不存在')
        require_project_access(project,u);bom=db.get(ProjectBomVersion,project.current_bom_version_id);template=db.get(ExcelTemplate,template_id)
        if not bom or not template:raise HTTPException(422,'请选择有效BOM和Excel模板')
        version=db.scalar(select(ExcelTemplateVersion).where(ExcelTemplateVersion.template_id==template.id,ExcelTemplateVersion.version==template.current_version));data,_,_=render_workbook(project,bom,version);digest=hashlib.sha256(data).hexdigest();file_name=f'{project.name}-BOM-V{bom.version}.xlsx';path=storage.put(f'exports/{uuid.uuid4().hex}-{file_name}',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');record=ExportRecord(project_id=project.id,bom_version_id=bom.id,template_version_id=version.id,export_by=u.id,file_path=path,file_name=file_name,file_hash=digest);db.add(record);db.flush();project.status='EXPORTED';audit(db,u,'EXPORT','project_bom',bom.id,{'recordId':record.id,'fileHash':digest});db.commit();return {'exportId':record.id,'fileName':file_name,'fileHash':digest,'downloadUrl':f'/api/exports/{record.id}/download'}

    @router.get('/api/exports/{export_id}/download')
    def download_export(export_id:int,db:Session=Depends(session),u=Depends(permit('BOM_EDIT'))):
        record=db.get(ExportRecord,export_id)
        if not record:raise HTTPException(404,'导出记录不存在')
        project=db.get(Project,record.project_id);require_project_access(project,u);response=storage.get(record.file_path.split('/',1)[1]);data=response.read();response.close();response.release_conn();return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f"attachment; filename*=UTF-8''{urllib.parse.quote(record.file_name)}",'X-File-SHA256':record.file_hash})

    @router.get('/api/documents/center/list')
    def document_center_list(status:str='',db:Session=Depends(session),u=Depends(permit('DOCUMENT_VIEW'))):
        query=select(DocumentAsset).order_by(DocumentAsset.updated_at.desc())
        if status:query=query.where(DocumentAsset.status==status)
        if 'DOCUMENT_EDIT' not in permission_codes(u):query=query.where(DocumentAsset.status=='PUBLISHED',DocumentAsset.document_status!='OBSOLETE')
        return [document_dto(d,u) for d in db.scalars(query)]

    @router.patch('/api/documents/{document_id}')
    def patch_document(document_id:int,x:DocumentPatch,db:Session=Depends(session),u=Depends(permit('DOCUMENT_EDIT'))):
        document=db.get(DocumentAsset,document_id)
        if not document:raise HTTPException(404,'资料不存在')
        values=x.model_dump(exclude_none=True)
        if values.get('status') not in (None,'DRAFT','PUBLISHED'):raise HTTPException(422,'发布状态无效')
        if values.get('document_status') not in (None,'CURRENT','HISTORICAL','OBSOLETE'):raise HTTPException(422,'资料有效性无效')
        if values.get('document_status')=='CURRENT':
            db.query(DocumentAsset).filter(DocumentAsset.id!=document.id,DocumentAsset.name==document.name,DocumentAsset.document_status=='CURRENT').update({'document_status':'HISTORICAL','knowledge_status':'OUTDATED'})
        for key,value in values.items():setattr(document,key,value)
        if not (document.status=='PUBLISHED' and document.document_status=='CURRENT' and document.knowledge_enabled):document.knowledge_status='OUTDATED' if document.knowledge_status=='SYNCED' else 'NOT_SYNCED'
        audit(db,u,'UPDATE','document',document.id,values);db.commit();return document_dto(document,u)

    @router.post('/api/documents/{document_id}/publish')
    def publish_document(document_id:int,db:Session=Depends(session),u=Depends(permit('DOCUMENT_EDIT'))):
        document=db.get(DocumentAsset,document_id)
        if not document:raise HTTPException(404,'资料不存在')
        document.status='PUBLISHED';document.knowledge_status='NOT_SYNCED' if document.knowledge_enabled and document.document_status=='CURRENT' else document.knowledge_status;audit(db,u,'PUBLISH','document',document.id);db.commit();return document_dto(document,u)

    @router.post('/api/documents/{document_id}/knowledge-sync')
    def sync_document(document_id:int,db:Session=Depends(session),u=Depends(permit('DOCUMENT_EDIT'))):
        document=db.get(DocumentAsset,document_id)
        if not document:raise HTTPException(404,'资料不存在')
        if not (document.status=='PUBLISHED' and document.knowledge_enabled and document.document_status=='CURRENT'):raise HTTPException(422,'只有已发布、当前有效且启用AI知识的资料可以同步')
        run=DocumentKnowledgeSync(document_id=document.id,status='PROCESSING',requested_by=u.id,started_at=datetime.now(timezone.utc));db.add(run);document.knowledge_status='PROCESSING';document.knowledge_sync_error='';db.commit();db.refresh(run)
        try:
            response=storage.get(document.path.split('/',1)[1]);content=response.read();response.close();response.release_conn()
            result=DifyKnowledgeClient().sync_file(file_name=document.original_file_name or document.name,mime_type=document.mime_type,content=content,existing_document_id=document.knowledge_document_id)
            finished=datetime.now(timezone.utc);run.status='SYNCED';run.finished_at=finished;run.dataset_id=result.dataset_id;run.external_document_id=result.document_id;document.knowledge_dataset_id=result.dataset_id;document.knowledge_document_id=result.document_id;document.knowledge_status='SYNCED';document.last_knowledge_sync_at=finished;document.knowledge_sync_error='';audit(db,u,'SYNC','document_knowledge',document.id,{'datasetId':result.dataset_id,'documentId':result.document_id,'batch':result.batch});db.commit();return {'status':'SYNCED','datasetId':result.dataset_id,'documentId':result.document_id,'batch':result.batch}
        except Exception as exc:
            message=' '.join(str(exc).replace('\x00',' ').split())[:800] or '未知错误';run.status='FAILED';run.finished_at=datetime.now(timezone.utc);run.error_message=message;document.knowledge_status='FAILED';document.knowledge_sync_error=message;audit(db,u,'SYNC_FAILED','document_knowledge',document.id,{'error':message});db.commit();raise HTTPException(502,f'知识同步失败：{message}') from exc

    return router
