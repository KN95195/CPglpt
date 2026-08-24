import html
import io
import os
from dataclasses import dataclass

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


class PreviewUnsupported(ValueError):
    pass


@dataclass(frozen=True)
class PreviewResult:
    preview_type: str
    data: bytes
    mime_type: str
    extension: str


def preview_capable(file_name: str, mime_type: str) -> bool:
    suffix = os.path.splitext(file_name or '')[1].lower()
    mime = (mime_type or '').lower()
    return bool(
        mime == 'application/pdf' or mime.startswith('image/') or mime.startswith('text/')
        or suffix in {'.pdf', '.png', '.jpg', '.jpeg', '.gif', '.webp', '.txt', '.docx', '.pptx', '.xlsx', '.xlsm'}
    )


def _text_pdf(pages: list[list[str]], *, wide: bool = False) -> bytes:
    output = io.BytesIO()
    page_size = landscape(A4) if wide else A4
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    pdf = canvas.Canvas(output, pagesize=page_size)
    width, height = page_size
    max_chars = 92 if wide else 62
    for page_number, lines in enumerate(pages or [['']], 1):
        y = height - 42
        pdf.setFont('STSong-Light', 9)
        pdf.drawRightString(width - 36, height - 24, f'{page_number}')
        for raw in lines:
            text = str(raw or '').replace('\t', '    ')
            chunks = [text[i:i + max_chars] for i in range(0, len(text), max_chars)] or ['']
            for chunk in chunks:
                if y < 42:
                    pdf.showPage()
                    pdf.setFont('STSong-Light', 9)
                    y = height - 42
                pdf.drawString(36, y, chunk)
                y -= 15
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def _docx_pdf(data: bytes) -> bytes:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(' | '.join(cell.text for cell in row.cells))
    return _text_pdf([lines])


def _pptx_pdf(data: bytes) -> bytes:
    presentation = Presentation(io.BytesIO(data))
    pages = []
    for index, slide in enumerate(presentation.slides, 1):
        lines = [f'第 {index} 页']
        for shape in slide.shapes:
            text = getattr(shape, 'text', '')
            if text:
                lines.extend(text.splitlines())
        pages.append(lines)
    return _text_pdf(pages, wide=True)


def xlsx_preview_html(data: bytes) -> bytes:
    workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    sections = []
    for sheet in workbook.worksheets:
        rows = []
        for values in sheet.iter_rows(max_row=min(sheet.max_row, 300), max_col=min(sheet.max_column, 60), values_only=True):
            cells = ''.join(f'<td>{html.escape(str(value if value is not None else ""))}</td>' for value in values)
            rows.append(f'<tr>{cells}</tr>')
        sections.append(f'<section><h2>{html.escape(sheet.title)}</h2><div class="sheet"><table>{"".join(rows)}</table></div></section>')
    content = '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>Excel预览</title><style>body{margin:20px;font:14px Arial,"Microsoft YaHei",sans-serif;color:#1f2937}h2{font-size:18px}.sheet{overflow:auto;border:1px solid #d8e0ea}table{border-collapse:collapse;white-space:nowrap}td{min-width:80px;padding:6px 8px;border:1px solid #d8e0ea;background:#fff}tr:first-child td{background:#eef5ff;font-weight:700}</style></head><body>'+''.join(sections)+'</body></html>'
    return content.encode('utf-8')


def build_preview(file_name: str, mime_type: str, data: bytes) -> PreviewResult | None:
    suffix = os.path.splitext(file_name or '')[1].lower()
    mime = (mime_type or '').lower()
    if mime == 'application/pdf' or mime.startswith('image/') or mime.startswith('text/') or suffix in {'.pdf', '.png', '.jpg', '.jpeg', '.gif', '.webp', '.txt'}:
        return None
    if suffix == '.docx':
        return PreviewResult('PDF', _docx_pdf(data), 'application/pdf', '.pdf')
    if suffix == '.pptx':
        return PreviewResult('PDF', _pptx_pdf(data), 'application/pdf', '.pdf')
    if suffix in {'.xlsx', '.xlsm'}:
        return PreviewResult('HTML', xlsx_preview_html(data), 'text/html; charset=utf-8', '.html')
    raise PreviewUnsupported('当前文件格式暂不支持在线预览，请下载原文件查看')
