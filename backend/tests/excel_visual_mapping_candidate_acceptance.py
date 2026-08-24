import io
import uuid

import httpx
import jwt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.config import settings
from app.database import SessionLocal
from app.models import ExcelTemplate, ExcelTemplateVersion, ExportRecord, Project, User
from app.storage import storage


def token_and_project():
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == 'admin'))
        project = db.scalar(select(Project).where(Project.current_bom_version_id.is_not(None)).order_by(Project.id))
        return jwt.encode({'uid': user.id}, settings.jwt_secret, algorithm='HS256'), project.id, project.name, project.customer


def template_bytes():
    workbook = Workbook()
    summary = workbook.active
    summary.title = '项目抬头'
    summary.merge_cells('A1:I1')
    summary['A1'] = '海智项目产品配单'
    summary['A1'].font = Font(bold=True, size=18, color='FFFFFF')
    summary['A1'].fill = PatternFill('solid', fgColor='146EF5')
    summary['A2'] = '项目名称'
    summary['E2'] = '客户名称'
    summary.freeze_panes = 'A3'
    bom = workbook.create_sheet('产品明细')
    bom.append(['标准BOM'])
    for _ in range(6):
        bom.append([])
    bom.append(['序号', '产品名称', '型号', '数量', '单位', '项目单价', '小计', '用途', '备注'])
    for cell in bom[8]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DCEBFF')
    bom.freeze_panes = 'A9'
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def main():
    token, project_id, project_name, customer_name = token_and_project()
    headers = {'Authorization': f'Bearer {token}'}
    template_id = export_id = None
    with httpx.Client(base_url='http://127.0.0.1:8000', headers=headers, timeout=90) as client:
        try:
            original = template_bytes()
            name = f'可视化映射验收-{uuid.uuid4().hex[:8]}'
            response = client.post('/api/excel-templates', files={'file': (name + '.xlsx', original, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, data={'name': name})
            assert response.status_code == 201, response.text
            template_id = response.json()['id']

            analysis = client.post(f'/api/excel-templates/{template_id}/analyze')
            assert analysis.status_code == 200, analysis.text
            payload = analysis.json()
            assert payload['sheetNames'] == ['项目抬头', '产品明细']
            assert {sheet['name'] for sheet in payload['sheets']} == {'项目抬头', '产品明细'}
            assert any(cell['coordinate'] == 'A1' for cell in payload['sheets'][0]['cells'])

            mapping = {
                'sheet_name': '产品明细',
                'placeholder_mappings': {'project_name': '项目抬头!B2', 'customer_name': '项目抬头!F2'},
                'bom_template_row': 8,
                'bom_column_mappings': {'lineNo': 'A', 'productName': 'B', 'model': 'C', 'quantity': 'D', 'unit': 'E', 'unitPrice': 'F', 'subtotal': 'G', 'purpose': 'H', 'note': 'I'},
            }
            saved = client.post(f'/api/excel-templates/{template_id}/mapping', json=mapping)
            assert saved.status_code == 200, saved.text

            preview = client.post(f'/api/projects/{project_id}/export-preview?template_id={template_id}')
            assert preview.status_code == 200, preview.text
            preview_payload = preview.json()
            assert preview_payload['readOnly'] is True
            assert project_name in preview_payload['html']
            assert '<table>' in preview_payload['html']

            exported = client.post(f'/api/projects/{project_id}/export-excel?template_id={template_id}')
            assert exported.status_code == 200, exported.text
            export_id = exported.json()['exportId']
            download = client.get(exported.json()['downloadUrl'])
            assert download.status_code == 200
            workbook = load_workbook(io.BytesIO(download.content), data_only=False)
            assert workbook.sheetnames == ['项目抬头', '产品明细']
            assert workbook['项目抬头']['B2'].value == project_name
            assert workbook['项目抬头']['F2'].value == customer_name
            assert 'A1:I1' in {str(item) for item in workbook['项目抬头'].merged_cells.ranges}
            assert workbook['项目抬头'].freeze_panes == 'A3'
            assert workbook['产品明细']['B8'].value
            assert workbook['产品明细']['A8'].fill.fgColor.rgb.endswith('DCEBFF')
            print('EXCEL_VISUAL_MAPPING_CANDIDATE_ACCEPTANCE_PASS', template_id, export_id)
        finally:
            with SessionLocal() as db:
                if export_id:
                    record = db.get(ExportRecord, export_id)
                    if record:
                        try:
                            storage.delete(record.file_path.split('/', 1)[1])
                        except Exception:
                            pass
                        db.delete(record)
                if template_id:
                    versions = list(db.scalars(select(ExcelTemplateVersion).where(ExcelTemplateVersion.template_id == template_id)))
                    version_ids = [version.id for version in versions]
                    for record in db.scalars(select(ExportRecord).where(ExportRecord.template_version_id.in_(version_ids))):
                        try:
                            storage.delete(record.file_path.split('/', 1)[1])
                        except Exception:
                            pass
                        db.delete(record)
                    db.flush()
                    for version in versions:
                        try:
                            storage.delete(version.file_path.split('/', 1)[1])
                        except Exception:
                            pass
                        db.delete(version)
                    db.flush()
                    template = db.get(ExcelTemplate, template_id)
                    if template:
                        db.delete(template)
                db.commit()


if __name__ == '__main__':
    main()
