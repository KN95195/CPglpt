import hashlib
import io
import json
import os
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone

import jwt
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas
from sqlalchemy import select

from app.config import settings
from app.database import SessionLocal
from app.models import User

BASE=os.environ.get('BASE_URL','http://127.0.0.1:18089')

def request(path,method='GET',body=None,token='',content_type='application/json'):
    data=None
    if body is not None:data=json.dumps(body,ensure_ascii=False).encode() if content_type=='application/json' else body
    headers={'Content-Type':content_type}
    if token:headers['Authorization']='Bearer '+token
    try:
        with urllib.request.urlopen(urllib.request.Request(BASE+path,data=data,headers=headers,method=method),timeout=40) as response:
            raw=response.read();return response.status,(json.loads(raw) if raw and 'json' in response.headers.get('Content-Type','') else raw)
    except urllib.error.HTTPError as exc:
        raw=exc.read()
        try:payload=json.loads(raw)
        except Exception:payload=raw.decode(errors='replace')
        return exc.code,payload

def multipart(fields,file_name,file_bytes,mime):
    boundary='----phase2acceptance'
    parts=[]
    for key,value in fields.items():parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode())
    parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{file_name}"\r\nContent-Type: {mime}\r\n\r\n'.encode()+file_bytes+b'\r\n')
    parts.append(f'--{boundary}--\r\n'.encode());return b''.join(parts),'multipart/form-data; boundary='+boundary

def login(username,password):
    status,payload=request('/api/auth/login','POST',{'username':username,'password':password})
    assert status==200,(username,status,payload);return payload['access_token']

def token_for(username):
    with SessionLocal() as db:
        user=db.scalar(select(User).where(User.username==username))
        assert user and user.enabled,username
        return jwt.encode({'uid':user.id,'exp':datetime.now(timezone.utc)+timedelta(minutes=30)},settings.jwt_secret,algorithm='HS256')

admin=token_for('admin')
sales=token_for('sales')
reader_username='phase2_reader_'+hashlib.sha256(os.urandom(16)).hexdigest()[:8]
reader_role='phase2_reader_role_'+reader_username[-8:]
reader_password='Phase2-Reader-Only-2026!'
status,_=request('/api/admin/roles','POST',{'code':reader_role,'name':'候选只读验收角色','permissions':['KNOWLEDGE_VIEW','DOCUMENT_VIEW','DOCUMENT_DOWNLOAD','TRAINING_VIEW']},admin);assert status==201
status,reader_user=request('/api/admin/users','POST',{'username':reader_username,'display_name':'候选只读验收用户','password':reader_password,'role_code':reader_role},admin);assert status==201
ordinary=login(reader_username,reader_password)
status,scenes=request('/api/scenes',token=sales);assert status==200 and scenes
status,products=request('/api/products',token=sales);assert status==200 and products
scene=next((row for row in scenes if row['name']=='桥梁防撞'),None)
assert scene, '候选数据库缺少桥梁防撞真实场景'

raw='某桥梁上下游各3公里，需要4个球机，需要AIS融合、船名OCR、偏航预警，7×24小时运行。'
status,parsed=request('/api/ai/config/parse','POST',{'raw_requirement':raw},sales);assert status==200 and parsed['parsedRequirement']['ptzCount']==4
status,project=request('/api/projects','POST',{'name':'桥梁防撞候选验收项目','customer':'候选环境客户','region':'浙江','scene_id':scene['id'],'raw_requirement':raw,'requirements':parsed['parsedRequirement']},sales);assert status==201
project_id=project['id']
status,confirmed=request('/api/ai/config/confirm','POST',{'project_id':project_id,'raw_requirement':raw,'ai_parsed_requirement':parsed['parsedRequirement'],'user_adjusted_requirement':parsed['parsedRequirement'],'confirmed_requirement':parsed['parsedRequirement']},sales);assert status==200
status,recommended=request('/api/bom/recommend','POST',{'project_id':project_id},sales);assert status==200 and recommended['items']
assert all(item.get('productId') and item.get('model') for item in recommended['items'])
original=json.loads(json.dumps(recommended['items'],ensure_ascii=False));edited=json.loads(json.dumps(original,ensure_ascii=False));edited[0]['quantity']=int(edited[0]['quantity'])+1;edited[0]['note']='人工验收调整'
status,version2=request(f'/api/projects/{project_id}/bom-versions','POST',{'items':edited,'change_summary':'候选验收人工调整','source':'MANUAL_EDIT'},sales);assert status==201 and version2['manualEdited'] and version2['version']==recommended['version']+1
status,versions=request(f'/api/projects/{project_id}/bom-versions',token=sales);assert status==200 and len(versions)>=2
old=next(v for v in versions if v['version']==recommended['version']);assert old['items'][0]['quantity']==original[0]['quantity']
status,validation=request('/api/bom/validate','POST',{'project_id':project_id},sales);assert status==200 and validation['status'] in {'PASS','WARNING'}

missing_product=next((p for p in products if p['id'] not in {i['productId'] for i in edited}),None)
if missing_product:
    status,rule=request('/api/bom/rules','POST',{'code':'PHASE2_ACCEPT_REQUIRE_'+reader_username[-8:].upper(),'name':'候选验收必选规则','rule_type':'REQUIRE','condition':{'sceneId':scene['id']},'action':{'productId':missing_product['id'],'quantity':1},'enabled':True},admin);assert status==201
    status,blocked=request('/api/bom/validate','POST',{'project_id':project_id,'override_reason':'销售不应绕过'},sales);assert status==403
    status,overridden=request('/api/bom/validate','POST',{'project_id':project_id,'override_reason':'候选环境产品经理审批验证'},admin);assert status==200 and overridden['overridden']

status,ordinary_project=request('/api/projects',token=ordinary);assert status==403
status,ordinary_ai=request('/api/ai/chat','POST',{'message':'海智AI终端参考价格是多少？'},ordinary);assert status==200 and not ordinary_ai['priceVisible']

workbook=Workbook();sheet=workbook.active;sheet.title='配单';sheet['A1']='海智项目配单';sheet['A2']='项目：{{project_name}}';sheet['E2']='客户：{{customer_name}}';headers=['序号','产品名称','型号','数量','单位','项目单价','小计','用途','备注']
for column,value in enumerate(headers,1):sheet.cell(7,column,value)
sheet.merge_cells('A1:I1');sheet.freeze_panes='A8';sheet.column_dimensions['B'].width=24;sheet.column_dimensions['C'].width=18
buffer=io.BytesIO();workbook.save(buffer);template_bytes=buffer.getvalue();body,mime=multipart({'name':'Phase 2候选验收模板'},'phase2-template.xlsx',template_bytes,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
status,template=request('/api/excel-templates','POST',body,admin,mime);assert status==201
template_id=template['id']
status,analysis=request(f'/api/excel-templates/{template_id}/analyze','POST',{},admin);assert status==200 and analysis['placeholders']['project_name']=='配单!A2'
columns={'lineNo':'A','productName':'B','model':'C','quantity':'D','unit':'E','unitPrice':'F','subtotal':'G','purpose':'H','note':'I'}
status,mapping=request(f'/api/excel-templates/{template_id}/mapping','POST',{'sheet_name':'配单','placeholder_mappings':analysis['placeholders'],'bom_template_row':8,'bom_column_mappings':columns},admin);assert status==200
status,preview=request(f'/api/projects/{project_id}/export-preview?template_id={template_id}','POST',{},sales);assert status==200 and len(preview['items'])==len(edited)
status,export=request(f'/api/projects/{project_id}/export-excel?template_id={template_id}','POST',{},sales);assert status==200
status,xlsx=request(export['downloadUrl'],token=sales);assert status==200 and hashlib.sha256(xlsx).hexdigest()==export['fileHash']
rendered=load_workbook(io.BytesIO(xlsx),data_only=False);assert rendered['配单']['B8'].value==edited[0]['productName'];assert rendered['配单']['D8'].value==edited[0]['quantity']

pdf_buffer=io.BytesIO();pdf_canvas=canvas.Canvas(pdf_buffer);pdf_canvas.drawString(72,760,'Phase 2 candidate knowledge acceptance');pdf_canvas.save();pdf=pdf_buffer.getvalue();body,mime=multipart({'category':'产品资料','version':'V1.0','description':'候选知识同步验收'},'phase2-acceptance.pdf',pdf,'application/pdf')
status,document=request('/api/documents','POST',body,admin,mime);assert status==201 and document['status']=='DRAFT'
document_id=document['id']
status,rejected=request(f'/api/documents/{document_id}/knowledge-sync','POST',{},admin);assert status==422
status,published=request(f'/api/documents/{document_id}','PATCH',{'status':'PUBLISHED','document_status':'CURRENT','knowledge_enabled':True},admin);assert status==200
status,synced=request(f'/api/documents/{document_id}/knowledge-sync','POST',{},admin);assert status==200 and synced['status']=='SYNCED'
status,downloaded=request(f'/api/documents/{document_id}/download',token=ordinary);assert status==200 and downloaded==pdf

status,_=request(f"/api/admin/users/{reader_user['id']}",'DELETE',token=admin);assert status==204
status,_=request(f'/api/admin/roles/{reader_role}','DELETE',token=admin);assert status==204
print(json.dumps({'marker':'PHASE2_CANDIDATE_API_ACCEPTANCE_PASS','projectId':project_id,'bomVersions':len(versions),'excelHash':export['fileHash'],'documentId':document_id,'migration':'c3d4e5f60718'},ensure_ascii=False))
