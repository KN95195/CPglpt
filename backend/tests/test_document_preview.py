import io
import unittest

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.document_preview import PreviewUnsupported, build_preview


class DocumentPreviewTest(unittest.TestCase):
    def test_docx_becomes_pdf(self):
        output = io.BytesIO()
        document = Document()
        document.add_heading('海智产品说明', 1)
        document.add_paragraph('桥梁防撞支持AIS融合、船名OCR和偏航预警。')
        document.save(output)
        preview = build_preview('说明书.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', output.getvalue())
        self.assertEqual(preview.mime_type, 'application/pdf')
        self.assertTrue(preview.data.startswith(b'%PDF-'))

    def test_pptx_becomes_pdf(self):
        output = io.BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = '桥梁防撞方案'
        slide.placeholders[1].text = 'AIS融合与偏航预警'
        presentation.save(output)
        preview = build_preview('方案.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', output.getvalue())
        self.assertEqual(preview.mime_type, 'application/pdf')
        self.assertTrue(preview.data.startswith(b'%PDF-'))

    def test_xlsx_becomes_html(self):
        output = io.BytesIO()
        workbook = Workbook()
        workbook.active.title = '标准BOM'
        workbook.active.append(['产品', '型号'])
        workbook.active.append(['海智AI分析终端', 'HZ-AI-01'])
        workbook.save(output)
        preview = build_preview('配单.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', output.getvalue())
        self.assertEqual(preview.mime_type, 'text/html; charset=utf-8')
        self.assertIn('海智AI分析终端', preview.data.decode())

    def test_pdf_is_direct_and_legacy_binary_is_rejected(self):
        self.assertIsNone(build_preview('资料.pdf', 'application/pdf', b'%PDF-1.4'))
        with self.assertRaises(PreviewUnsupported):
            build_preview('旧资料.doc', 'application/msword', b'binary')


if __name__ == '__main__':
    unittest.main()
