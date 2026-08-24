import hashlib
import io

import httpx
import jwt
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from sqlalchemy import select

from app.config import settings
from app.database import SessionLocal
from app.models import User


def admin_token():
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == 'admin'))
        return jwt.encode({'uid': user.id}, settings.jwt_secret, algorithm='HS256')


def samples():
    output = io.BytesIO()
    document = Document()
    document.add_heading('海智产品说明', 1)
    document.add_paragraph('桥梁防撞支持AIS融合、船名OCR和偏航预警。')
    document.save(output)
    yield '海智说明书.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', output.getvalue(), 'application/pdf', b'%PDF-'

    output = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = '海智桥梁防撞方案'
    slide.placeholders[1].text = 'AIS融合和偏航预警'
    presentation.save(output)
    yield '海智方案.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', output.getvalue(), 'application/pdf', b'%PDF-'

    output = io.BytesIO()
    workbook = Workbook()
    workbook.active.title = '标准BOM'
    workbook.active.append(['产品', '型号'])
    workbook.active.append(['海智AI分析终端', 'HZ-AI-01'])
    workbook.save(output)
    yield '海智配单.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', output.getvalue(), 'text/html; charset=utf-8', '海智AI分析终端'.encode()


def main():
    headers = {'Authorization': f'Bearer {admin_token()}'}
    created = []
    with httpx.Client(base_url='http://127.0.0.1:8000', headers=headers, timeout=90) as client:
        try:
            for name, mime, original, preview_mime, marker in samples():
                response = client.post('/api/documents', files={'file': (name, original, mime)}, data={'description': '候选预览验收'})
                assert response.status_code == 201, (name, response.status_code, response.text)
                payload = response.json()
                created.append(payload['id'])
                assert payload['previewStatus'] == 'READY', payload

                preview = client.get(f"/api/documents/{payload['id']}/preview")
                assert preview.status_code == 200, (name, preview.status_code, preview.text)
                assert preview.headers['content-type'].startswith(preview_mime.split(';')[0])
                assert marker in preview.content

                download = client.get(f"/api/documents/{payload['id']}/download")
                assert download.status_code == 200
                assert hashlib.sha256(download.content).hexdigest() == hashlib.sha256(original).hexdigest()
            print('DOCUMENT_PREVIEW_CANDIDATE_ACCEPTANCE_PASS', created)
        finally:
            for document_id in created:
                response = client.delete(f'/api/documents/{document_id}')
                assert response.status_code == 204, (document_id, response.status_code, response.text)


if __name__ == '__main__':
    main()
