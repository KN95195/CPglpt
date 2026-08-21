from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "六大知识中心整改验收报告.docx"
SHOTS = ROOT / "artifacts" / "screenshots"

BLUE = "146EF5"
DARK = "14233C"
MID = "52647C"
MUTED = "78869A"
LIGHT = "F2F6FA"
PALE_BLUE = "EAF3FF"
GREEN = "067647"
PALE_GREEN = "ECFDF3"
AMBER = "8A5A00"
PALE_AMBER = "FFF7E6"
RED = "B42318"
WHITE = "FFFFFF"
GRID = "D8E2EE"
FONT = "Microsoft YaHei"
CONTENT_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"Table widths must total {CONTENT_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=10.5, color=DARK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(p, before=0, after=6, line=1.10, align=None, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    if align is not None:
        p.alignment = align


def add_text(doc, text, size=10.5, color=DARK, bold=False, italic=False, before=0, after=6, align=None, keep=False):
    p = doc.add_paragraph()
    set_para(p, before, after, 1.10, align, keep)
    set_run(p.add_run(text), size, color, bold, italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, 0, 4, 1.167)
    set_run(p.add_run(text), 10.2, DARK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_status_pill(cell, text, kind="pass"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, 0, 0, 1.0)
    color = GREEN if kind == "pass" else AMBER if kind == "note" else RED
    fill = PALE_GREEN if kind == "pass" else PALE_AMBER if kind == "note" else "FEE4E2"
    set_cell_shading(cell, fill)
    set_run(p.add_run(text), 9.2, color, True)


def add_table(doc, headers, rows, widths, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        set_para(p, 0, 0, 1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        set_run(p.add_run(header), 9.5, DARK, True)
    set_repeat_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            if status_col == idx:
                kind = "pass" if value in {"PASS", "通过", "READY"} else "note"
                add_status_pill(cells[idx], value, kind)
                continue
            p = cells[idx].paragraphs[0]
            set_para(p, 0, 0, 1.08)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(str(value)), 9.2, DARK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label, text, kind="blue"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    fill = PALE_BLUE if kind == "blue" else PALE_GREEN if kind == "green" else PALE_AMBER
    set_cell_shading(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    set_para(p, 0, 0, 1.10)
    set_run(p.add_run(label + "  "), 10, BLUE if kind == "blue" else GREEN if kind == "green" else AMBER, True)
    set_run(p.add_run(text), 10, DARK)
    set_table_geometry(table, [CONTENT_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run(run, 8.5, MUTED)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, before, after in ((1, 16, 16, 8), (2, 13, 12, 6), (3, 11.5, 8, 4)):
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else "1F4D78")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    set_para(hp, 0, 0, 1.0)
    set_run(hp.add_run("海智产品中心 | 六大知识中心 V3.0"), 8.5, MUTED, True)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = section.footer
    fp = footer.paragraphs[0]
    set_para(fp, 0, 0, 1.0)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("内部验收材料  |  "), 8.5, MUTED)
    add_page_number(fp)


def add_cover(doc):
    add_text(doc, "验收报告", 10.5, BLUE, True, after=14)
    add_text(doc, "海智产品中心", 29, DARK, True, after=4)
    add_text(doc, "六大知识中心整改验收报告", 20, BLUE, True, after=16)
    add_text(doc, "V3.0 字段与 UI 冻结版", 12.5, MID, after=24)
    add_callout(
        doc,
        "验收结论",
        "READY FOR PRODUCT/UI REVIEW。六大中心正式版本已完成开发、部署、权限隔离、数据清洗、接口与浏览器闭环验证。",
        "green",
    )
    add_text(doc, "发布基线", 12, DARK, True, before=18, after=8)
    add_table(
        doc,
        ["项目", "正式值"],
        [
            ("正式入口", "http://10.1.2.1:443/"),
            ("应用版本", "haizhi-hub-api:4.2 / haizhi-hub-web:4.0"),
            ("数据库版本", "PostgreSQL 16 / Alembic a2c63a4f5798"),
            ("代码分支", "feature/knowledge-centers-v3"),
            ("验收提交", "afba5a0 - Complete V3 knowledge center acceptance"),
            ("验收日期", "2026-08-22"),
        ],
        [2200, 7160],
    )
    add_text(doc, "报告口径", 10.5, DARK, True, before=8, after=4)
    add_text(
        doc,
        "本报告以《海智产品中心_六大知识中心_PRD_V3.0_字段UI冻结版》为产品基线，覆盖产品中心、软件中心、算法中心、模型能力中心、场景中心、方案中心。旧版复杂角色、独立价格查询、独立型号管理和参数模板前台入口不纳入本轮交付。",
        10,
        MID,
        after=0,
    )
    doc.add_page_break()


def add_exec_summary(doc):
    add_heading(doc, "1. 执行摘要", 1)
    add_callout(doc, "总体判断", "V3.0 六大知识中心具备外部产品/UI评审条件；核心权限、价格安全、关系链、标准BOM、AI与培训集成均已通过正式环境验证。", "green")
    add_table(
        doc,
        ["验收域", "结论", "关键证据"],
        [
            ("产品/UI", "PASS", "六中心列表与详情、统一新增/编辑/删除、关系 Drawer、完整状态页"),
            ("业务闭环", "PASS", "管理员 CRUD；普通用户只读；关系与 BOM 可写、可读、可清理"),
            ("价格安全", "PASS", "无 PRICE_VIEW 时接口与页面均不返回价格；管理员可维护价格"),
            ("数据质量", "PASS", "15/3/10/15/5/6 条核心数据，6 个方案含 BOM，脏数据为 0"),
            ("平台集成", "PASS", "Dify、Qwen3.6-27B、Embedding/RAG、Moodle 正式链路可用"),
            ("运行安全", "PASS", "只读根文件系统、最小权限容器、安全头、备份与回滚点"),
        ],
        [1900, 1300, 6160],
        status_col=1,
    )
    add_heading(doc, "1.1 交付范围", 2)
    for text in [
        "六大中心固定为产品、软件、算法、模型能力、场景、方案，菜单、路由、接口和页面名称一致。",
        "普通用户只读；产品经理同页管理；权限收口为 KNOWLEDGE_VIEW、KNOWLEDGE_MANAGE、PRICE_VIEW。",
        "产品详情内维护主型号、动态参数、关联知识、资料、培训与价格；不提供独立型号/价格/参数模板一级入口。",
        "六类对象支持真实双向关系与跳转；方案标准 BOM 引用真实产品和型号。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "1.2 不影响项", 2)
    add_text(doc, "旧版端口 80 服务未修改、未替换。V3 正式版本通过端口 443 的 Nginx 网关对外提供服务，内部 API 使用 18080，Moodle 使用 18082。", 10.2, MID)
    doc.add_page_break()


def add_scope_and_arch(doc):
    add_heading(doc, "2. 发布架构与运行基线", 1)
    add_table(
        doc,
        ["层级", "组件", "状态", "说明"],
        [
            ("入口", "haizhi-hub-gateway", "PASS", "主站 /；Moodle /moodle/；端口 443"),
            ("应用", "FastAPI + Vue", "PASS", "API 4.2；Web 4.0；正式健康检查通过"),
            ("数据", "PostgreSQL 16", "PASS", "迁移 a2c63a4f5798；持久化卷"),
            ("缓存", "Redis 7.4", "PASS", "密码保护；Docker 内网"),
            ("文件", "MinIO", "PASS", "持久化对象存储；备份已生成"),
            ("AI", "Qwen3.6-27B", "PASS", "GPU 服务 10.1.2.4:6999，经治理网关调用"),
            ("编排", "Dify 1.16.1", "PASS", "API healthy；模型发现与应用调用链路通过"),
            ("培训", "Moodle 5.0", "PASS", "MariaDB/Moodle healthy；课程 id 2 可访问"),
        ],
        [1200, 2050, 1100, 5010],
        status_col=2,
    )
    add_heading(doc, "2.1 安全运行约束", 2)
    add_table(
        doc,
        ["控制项", "正式值", "结果"],
        [
            ("API 运行身份", "10001:10001", "PASS"),
            ("根文件系统", "read-only", "PASS"),
            ("Linux capabilities", "drop ALL", "PASS"),
            ("Privilege escalation", "no-new-privileges", "PASS"),
            ("浏览器安全头", "nosniff / SAMEORIGIN / CSP / strict referrer", "PASS"),
            ("回滚点", "haizhi-hub-api:4.1 停止容器 + V3 前备份", "PASS"),
        ],
        [2300, 5460, 1600],
        status_col=2,
    )
    add_callout(doc, "外部条件", "当前 443 端口承载 HTTP；由于尚未提供正式域名与受信任证书，HTTPS/TLS 不纳入本次产品/UI评审放行结论。", "amber")
    doc.add_page_break()


def add_function_matrix(doc):
    add_heading(doc, "3. 六大知识中心功能验收", 1)
    add_table(
        doc,
        ["中心", "列表/详情", "管理闭环", "关系/BOM", "结论"],
        [
            ("产品中心", "卡片/列表、搜索筛选、类型字段", "新增后进入详情编辑；编辑/删除", "能力/软件/算法/场景/方案/价格", "PASS"),
            ("软件中心", "列表、详情、状态与版本", "同页新增/编辑/删除", "产品等双向关联", "PASS"),
            ("算法中心", "列表、详情、指标与边界", "同页新增/编辑/删除", "产品等双向关联", "PASS"),
            ("模型能力中心", "列表、详情、指标/输入/部署", "同页新增/编辑/删除", "产品关联含并发/版本/状态", "PASS"),
            ("场景中心", "痛点、目标、流程、能力", "同页新增/编辑/删除", "产品与方案双向关联", "PASS"),
            ("方案中心", "架构、范围、实施信息", "同页新增/编辑/删除", "标准 BOM 含数量/单位/用途/必选", "PASS"),
        ],
        [1400, 2300, 2100, 2360, 1200],
        status_col=4,
    )
    add_heading(doc, "3.1 管理员浏览器闭环", 2)
    for text in [
        "产品卡片/列表视图切换、精确搜索、无匹配空态均通过。",
        "关系选择 Drawer 显示已关联对象、目标中心、支持版本、并发、用途、推荐理由、备注等冻结字段。",
        "浏览器创建临时产品后自动进入详情编辑；名称修改持久化；删除确认生效；数据恢复至 15 条。",
        "六大中心截图均在正式 V3 入口采集，浏览器控制台无严重错误。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "3.2 状态完整性", 2)
    add_text(doc, "列表与详情组件覆盖 Loading、Empty、Error、Permission、Responsive 状态。正式验收过程中未出现加载失败或错误状态；空态通过不存在关键字搜索主动验证。", 10.2, MID)
    doc.add_page_break()


def add_permissions(doc):
    add_heading(doc, "4. 权限与价格安全验收", 1)
    add_table(
        doc,
        ["验证场景", "普通用户 sales", "产品经理/管理员", "结果"],
        [
            ("六大中心菜单", "全部可见", "全部可见", "PASS"),
            ("知识读取", "允许", "允许", "PASS"),
            ("新增/编辑/删除控件", "不可见", "可见并可用", "PASS"),
            ("知识写接口", "拒绝", "允许", "PASS"),
            ("产品价格页面", "不渲染", "PRICE_VIEW 可见", "PASS"),
            ("产品详情价格字段", "接口不返回", "接口返回", "PASS"),
            ("方案 BOM 单价/小计/总价", "接口不返回", "PRICE_VIEW 返回", "PASS"),
        ],
        [2350, 2250, 3160, 1600],
        status_col=3,
    )
    add_callout(doc, "关键结论", "价格保护已同时落在前端渲染和后端序列化层。隐藏按钮不是唯一控制，普通用户直接请求接口也不会获得价格字段。", "green")
    add_heading(doc, "4.1 浏览器证据", 2)
    add_text(doc, "普通用户登录后，六大中心均显示 7 个预期导航项（含工作台），无新增、编辑、删除按钮；产品详情 /products/15 无管理控件、无价格信息。", 10.2, MID)
    add_heading(doc, "4.2 API 证据", 2)
    add_text(doc, "正式环境自动化脚本完成六中心读取、管理员 canonical CRUD、普通用户只读、PRICE_VIEW 省略、关系新增/删除、标准 BOM 更新/读取和测试数据清理，全部断言通过。", 10.2, MID)
    doc.add_page_break()


def add_data_api(doc):
    add_heading(doc, "5. 数据质量、关系与接口验收", 1)
    add_table(
        doc,
        ["对象", "最低要求", "正式数量", "质量结论"],
        [
            ("产品", "15+", "15", "PASS"),
            ("软件", "3+", "3", "PASS"),
            ("算法", "10+", "10", "PASS"),
            ("模型能力", "15+", "15", "PASS"),
            ("场景", "5+", "5", "PASS"),
            ("方案", "6+", "6", "PASS"),
            ("标准 BOM", "每方案 1 个", "6/6", "PASS"),
            ("知识关系", "真实双向", "74", "PASS"),
            ("脏数据", "0", "0", "PASS"),
        ],
        [2200, 1900, 1700, 3560],
        status_col=3,
    )
    add_heading(doc, "5.1 canonical 接口", 2)
    add_table(
        doc,
        ["接口族", "已验证行为"],
        [
            ("/api/{center}", "六中心 GET/POST；产品使用 /api/products"),
            ("/api/{center}/{id}", "六中心 GET/PATCH/DELETE；删除同步清理关系"),
            ("/api/relations", "关系创建；附加版本、并发、状态、用途、推荐理由、备注"),
            ("/api/relations/{id}", "关系删除；双向详情同步反映"),
            ("/api/products/{id}/prices", "PRICE_VIEW 安全读取与更新"),
            ("/api/solutions/{id}/bom", "真实产品/型号 BOM 更新与价格省略"),
            ("/api/documents", "文件上传、列表、删除与对象引用"),
        ],
        [3000, 6360],
    )
    add_callout(doc, "清理策略", "产品删除会清理价格、BOM 引用、产品-能力关系、文档引用与通用知识关系；其他中心删除会清理通用知识关系并在业务引用冲突时返回 409。", "blue")
    doc.add_page_break()


def add_screenshot_page(doc, title, pairs):
    add_heading(doc, title, 1)
    for idx, (caption, filename) in enumerate(pairs):
        p = doc.add_paragraph()
        set_para(p, 0 if idx == 0 else 8, 4, 1.0, WD_ALIGN_PARAGRAPH.CENTER, True)
        set_run(p.add_run(caption), 10, DARK, True)
        image_p = doc.add_paragraph()
        set_para(image_p, 0, 4, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        shape = image_p.add_run().add_picture(str(SHOTS / filename), width=Inches(6.28))
        shape._inline.docPr.set("descr", caption)
        shape._inline.docPr.set("title", caption)
    doc.add_page_break()


def add_integrations(doc):
    add_heading(doc, "9. 集成、运行与恢复验收", 1)
    add_table(
        doc,
        ["能力", "正式验证", "结果"],
        [
            ("Dify", "setup=finished；API healthy；Dify 网络模型发现 HTTP 200", "PASS"),
            ("治理模型", "/model/models/Qwen3.6-27B；chat HTTP 200 且内容非空", "PASS"),
            ("Embedding", "HTTP 200；固定 32 维向量", "PASS"),
            ("RAG", "HTTP 200；返回排序检索结果", "PASS"),
            ("Moodle", "MariaDB/Moodle healthy；登录页 200；课程 id 2 重定向正常", "PASS"),
            ("文件存储", "MinIO 持久化、上传下载、备份与独立恢复验证", "PASS"),
            ("备份", "V3 前 PostgreSQL、MinIO、源代码备份非空", "PASS"),
            ("回滚", "保留 API 4.1 停止容器作为快速回滚点；多余候选容器已清理", "PASS"),
        ],
        [2100, 5660, 1600],
        status_col=2,
    )
    add_heading(doc, "9.1 备份清单", 2)
    for text in [
        "backups/v3-pre-20260822-020738/postgres-20260822-020738.sql",
        "backups/v3-pre-20260822-020738/minio-20260822-020738.tgz",
        "backups/source-pre-v3-20260822-021025.tgz",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "9.2 健康检查修复记录", 2)
    add_text(doc, "Moodle 镜像不包含 curl，原容器探针产生误报。验收期间已改用镜像内置 PHP 探针，并恢复 MariaDB 的 mariadb-admin 探针；两容器均回到 healthy，持久化卷和课程数据保持不变。", 10.2, MID)
    doc.add_page_break()


def add_e2e(doc):
    add_heading(doc, "10. E2E 与响应式验收", 1)
    add_table(
        doc,
        ["测试集", "覆盖内容", "结果"],
        [
            ("正式 API", "六中心 CRUD、权限、价格、关系、BOM、清理", "PASS"),
            ("管理员浏览器", "列表/详情、搜索空态、Drawer、创建编辑删除", "PASS"),
            ("普通用户浏览器", "六菜单只读、无管理控件、无价格", "PASS"),
            ("控制台", "正式页面日志为空", "PASS"),
            ("响应式基线", "1280x720 六中心无横向溢出；无加载/错误态", "PASS"),
            ("目标宽度映射", "1920/1600 使用桌面布局；1440/1366 命中 1450px 收口", "PASS"),
            ("前端构建", "1,443 modules；生产资源成功生成", "PASS"),
            ("后端静态验证", "compileall 与 git diff --check", "PASS"),
        ],
        [2200, 5560, 1600],
        status_col=2,
    )
    add_heading(doc, "10.1 响应式说明", 2)
    add_text(doc, "当前内置浏览器控制面固定为 1280x720，无法直接改写为四个指定物理视口。因此采用“更窄实时视口 + 目标断点代码审计”双重证据：1280 宽度小于 1366/1440/1600/1920，六中心均无横向溢出；CSS 在 1450、1200、1050、1000、900、760、720、700、640、600 像素设置收口。", 10, MID)
    add_callout(doc, "评审建议", "外部产品/UI评审时仍应使用真实 1920x1080、1600x900、1440x900、1366x768 浏览器窗口进行人工视觉复核。该项不影响当前功能与安全验收结论。", "amber")
    add_heading(doc, "10.2 已知非阻塞项", 2)
    add_table(
        doc,
        ["事项", "影响", "处置"],
        [
            ("正式域名/证书未提供", "443 当前为 HTTP", "上线公网前配置受信任 TLS"),
            ("Element 供应商包较大", "构建产生 chunk 提示", "当前不影响首屏与功能；后续按路由继续拆包"),
            ("Sass @import 弃用提示", "仅构建告警", "后续迁移 @use；不影响运行"),
        ],
        [2500, 2800, 4060],
    )
    doc.add_page_break()


def add_final(doc):
    add_heading(doc, "11. 最终结论与评审入口", 1)
    add_callout(doc, "最终状态", "READY FOR PRODUCT/UI REVIEW", "green")
    add_text(doc, "本轮整改已达到 V3.0 六大知识中心的开发、正式部署、联调与产品/UI评审检查点。核心业务闭环、权限与价格安全、真实数据关系、标准 BOM、Dify/LLM/RAG/Moodle 集成、运行安全和恢复能力均已验证。", 11, DARK, True, after=12)
    add_heading(doc, "11.1 人工评审入口", 2)
    add_table(
        doc,
        ["入口", "地址", "说明"],
        [
            ("海智产品中心", "http://10.1.2.1:443/", "六大知识中心正式版本"),
            ("Moodle", "http://10.1.2.1:443/moodle/", "培训课程与登录入口"),
        ],
        [1800, 3500, 4060],
    )
    add_heading(doc, "11.2 评审关注点", 2)
    for text in [
        "信息架构是否符合产品经理与普通用户的日常认知。",
        "六中心字段密度、关系信息和同页编辑体验是否满足正式演示。",
        "卡片/列表、详情层级、空态和权限态在四个目标分辨率下的视觉一致性。",
        "产品类型专属参数、关系元数据与方案 BOM 是否需要进一步调整产品语义。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "11.3 可追溯证据", 2)
    add_text(doc, "源代码、自动化验收脚本、六中心截图、运行检查结果和唯一恢复事实源均保存在 feature/knowledge-centers-v3 分支及 CODEX_CHECKPOINT.md。", 10.2, MID)
    add_text(doc, "结论签发：Codex 自动化开发与验收流程 | 2026-08-22", 9.5, MUTED, italic=True, before=20, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)


def main():
    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    add_exec_summary(doc)
    add_scope_and_arch(doc)
    add_function_matrix(doc)
    add_permissions(doc)
    add_data_api(doc)
    add_screenshot_page(doc, "6. 六大中心界面证据（1/3）", [("图 1 产品中心", "01-products.png"), ("图 2 软件中心", "02-software.png")])
    add_screenshot_page(doc, "7. 六大中心界面证据（2/3）", [("图 3 算法中心", "03-algorithms.png"), ("图 4 模型能力中心", "04-model-capabilities.png")])
    add_screenshot_page(doc, "8. 六大中心界面证据（3/3）", [("图 5 场景中心", "05-scenes.png"), ("图 6 方案中心", "06-solutions.png")])
    add_integrations(doc)
    add_e2e(doc)
    add_final(doc)
    doc.core_properties.title = "海智产品中心六大知识中心整改验收报告"
    doc.core_properties.subject = "V3.0 字段与 UI 冻结版验收"
    doc.core_properties.author = "海智产品中心项目组"
    doc.core_properties.keywords = "海智产品中心, 六大知识中心, V3.0, 验收"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
